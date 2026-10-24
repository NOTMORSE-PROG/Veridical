"""F3.5 fake-mode quotes are source text, never source-coverage metadata."""

import json
from copy import deepcopy

import pytest
from docx import Document

from app.checks.semantic import _verify_quotes
from app.ingest.analytical_source import analytical_source
from app.ingest.schemas import TextBlock
from app.llm.fake import FIXTURES_DIR, FakeLLMClient, _real_quotes_from_prompt, _with_real_quotes
from tests.test_native_math_source import _extract, _math


def prompt(text):
    return f"<manuscript_text>\n{text}\n</manuscript_text>"


def envelope(scopes):
    return json.dumps(
        {
            "source_coverage": {"native_math_coverage_version": "v1"},
            "source_limitation": "Coverage metadata with many words is never source evidence.",
            "contiguous_source_scopes": scopes,
        }
    )


async def test_actual_native_source_envelope_yields_only_verifiable_block_quotes(tmp_path):
    document = Document()
    document.add_paragraph(
        "Readable alpha and beta remain literal source text throughout this synthetic control."
    )
    mixed = document.add_paragraph("Affected content must not enter an evidence quote.")
    _math(mixed)
    document.add_paragraph("A separate readable source block supports only its own recorded words.")
    extraction = _extract(tmp_path, document)
    source = analytical_source(extraction)
    text = source.context(extraction.blocks)
    assert json.loads(text)["source_coverage"]["n_native_equations_unassessed"] == 1
    quotes = _real_quotes_from_prompt(prompt(text))
    assert len(quotes) == 2
    assert all(
        _verify_quotes([quote], extraction.blocks, extraction.anchor_kind, scopes=source.scopes)
        is not None
        for quote in quotes
    )
    assert all(
        "source_coverage" not in quote and "Affected content" not in quote for quote in quotes
    )
    scripted = json.loads((FIXTURES_DIR / "semantic_grading.json").read_text(encoding="utf-8"))
    response = await FakeLLMClient().complete("semantic_grading", prompt(text))
    assert [
        {k: v for k, v in item.items() if k != "evidence_quotes"} for item in response["verdicts"]
    ] == [
        {k: v for k, v in item.items() if k != "evidence_quotes"} for item in scripted["verdicts"]
    ]


@pytest.mark.parametrize(
    "text",
    [
        envelope([]),
        "[" * 1200 + "0" + "]" * 1200,
        '{"source_coverage": this JSON is malformed and must never be quoted as prose',
        envelope([{"text": "A wrongly nested block contains many tempting words for a quote."}]),
        envelope([[{"paragraph": 0, "text": True}]]),
        envelope(
            [
                [
                    {"paragraph": 0, "text": "Only four source words"},
                    {"paragraph": 1, "text": "Another short source fragment"},
                ]
            ]
        ),
        envelope(
            [
                [{"paragraph": 0, "text": "Only four source words"}],
                [{"paragraph": 1, "text": "Another short source fragment"}],
            ]
        ),
        "No manuscript wrapper",
    ],
)
def test_empty_or_invalid_source_never_invents_or_joins_quote_candidates(text):
    original = {
        "verdicts": [
            {
                "index": 0,
                "verdict": "pass",
                "reasoning": "Scripted fixture only",
                "evidence_quotes": ["Canned placeholder never present in the source"],
            }
        ],
        "extra": "preserved",
    }
    before = deepcopy(original)
    rendered = text if text == "No manuscript wrapper" else prompt(text)
    assert _real_quotes_from_prompt(rendered) == []
    result = _with_real_quotes(original, rendered)
    assert result["verdicts"] == [{**original["verdicts"][0], "evidence_quotes": []}]
    assert result["extra"] == "preserved" and original == before


def test_plain_source_and_decoded_unicode_lines_keep_their_existing_boundaries():
    first = "Readable α and β remain source words without invented transcriptions."
    second = "A separate line has enough actual manuscript words for this control."
    assert _real_quotes_from_prompt(prompt(first + "\n" + second)) == [
        " ".join(first.split()[:8]),
        " ".join(second.split()[:8]),
    ]
    structured = envelope([[{"paragraph": 0, "text": first + "\n" + second}]])
    assert _real_quotes_from_prompt(prompt(structured)) == _real_quotes_from_prompt(
        prompt(first + "\n" + second)
    )


def test_valid_json_without_source_envelope_keys_remains_literal_manuscript_text():
    text = json.dumps(
        {"example": "A legitimate JSON code example contains enough quoted source words."}
    )
    blocks = [TextBlock(paragraph=0, text=text, max_font_size=12, bold_ratio=0)]
    quotes = _real_quotes_from_prompt(prompt(text))
    assert quotes and quotes[0].startswith('{"example":')
    assert _verify_quotes(quotes, blocks, "paragraph") is not None
