"""F5 source membership survives ingestion, persistence and separate F3 review."""

import hashlib
import os
from decimal import Decimal
from unittest.mock import AsyncMock

import pytest
from docx import Document
from sqlalchemy import delete, select

from app.checks.citations import verify
from app.checks.router import route_criterion
from app.checks.rules.context import build_rule_context
from app.checks.structural import run_structural_check
from app.config import Settings
from app.external.schemas import VerificationResult
from app.ingest.patterns import load_patterns
from app.ingest.service import ingest_manuscript, load_raw_store_async
from app.models.audit import AuditLog
from app.models.citation import Citation
from app.models.enums import CheckRunStatus, CriterionType, ResultOutcome
from app.models.instructor import Instructor
from app.models.manuscript import Manuscript
from app.models.rubric import Criterion, Rubric
from app.models.run import CheckResult, CheckRun, Flag
from app.report.service import aggregate_and_score, get_report, resolve_escalation_for_run
from tests.test_citation_claim_occurrences import ScriptedClaims
from tests.test_reference_inventory import REF_A, REF_B
from tests.test_report_get_live import _clean_tables as _clean_tables
from tests.test_report_get_live import scratch_url as scratch_url
from tests.test_report_get_live import session_factory as session_factory

pytestmark = pytest.mark.skipif(
    "DATABASE_URL" not in os.environ, reason="requires private Postgres"
)


@pytest.mark.parametrize("partial", [False, True])
async def test_fresh_f5_inventory_boundaries_survive_separate_criterion_review(
    session_factory, tmp_path, monkeypatch, partial
):
    document = Document()
    document.add_paragraph("Prior work (Reyes, 2024) and (Santos, 2022) found improvements.")
    document.add_heading("References", 1)
    document.add_paragraph(REF_A)
    document.add_paragraph(REF_B)
    document.add_heading("Appendix", 1)
    path = tmp_path / "source.docx"
    document.save(path)
    original = path.read_bytes()
    settings = Settings(_env_file=None, data_dir=tmp_path / "data")
    metadata = AsyncMock(
        return_value=verify.CitationVerdict(
            verify.VerdictKind.existence_confirmed,
            VerificationResult(
                found=True, provider="crossref", abstract="The study found improvements."
            ),
        )
    )
    monkeypatch.setattr(verify, "verify_citation", metadata)
    model = ScriptedClaims()
    async with session_factory() as session:
        owner = Instructor(email="inventory@example.test", display_name="Inventory control")
        session.add(owner)
        await session.commit()
        manuscript = Manuscript(
            instructor_id=owner.id,
            group_label="Synthetic",
            file_ref=str(path),
            content_hash=hashlib.sha256(original).hexdigest(),
        )
        rubric = Rubric(instructor_id=owner.id, title="Synthetic", source_file="format.docx")
        session.add_all([manuscript, rubric])
        await session.commit()
        await ingest_manuscript(session, manuscript, path, settings, read_images=False)
        criterion = Criterion(
            rubric_id=rubric.id,
            type=CriterionType.structural,
            text="At least 2 references",
            weight=Decimal("1"),
            position=0,
        )
        session.add(criterion)
        await session.commit()
        if partial:
            await session.execute(
                delete(Citation).where(
                    Citation.manuscript_id == manuscript.id, Citation.order_index == 1
                )
            )
            await session.commit()
        run = CheckRun(manuscript_id=manuscript.id, rubric_id=rubric.id, status=CheckRunStatus.done)
        session.add(run)
        await session.commit()
        ctx = await build_rule_context(session, manuscript, settings)
        extraction = await load_raw_store_async(settings, manuscript.id)
        source_before = extraction.model_dump_json()
        citations = list(
            await session.scalars(
                select(Citation)
                .where(Citation.manuscript_id == manuscript.id)
                .order_by(Citation.order_index)
            )
        )
        routed = route_criterion(criterion, criterion_id=criterion.id, raw_type="structural")
        structural = await run_structural_check(
            session, run.id, criterion, criterion.id, routed, ctx
        )
        result = await verify.run_citation_integrity_check(
            session, None, run.id, citations, extraction, load_patterns(), settings, model
        )
        assert extraction.model_dump_json() == source_before
        assert not list(
            await session.scalars(select(Flag).where(Flag.check_result_id == result.id))
        )
        await aggregate_and_score(session, run.id, settings, commit=True)
        report = await get_report(session, run.id, owner.id)
        if partial:
            assert structural.outcome == ResultOutcome.unverifiable
            assert result.outcome == ResultOutcome.escalated
            assert report.status == "needs_review" and report.pending_review_count == 1
            resolved = await resolve_escalation_for_run(
                session,
                run.id,
                structural.id,
                owner.id,
                "mark_pass",
                "Scripted manual criterion override; the reference inventory remains unavailable.",
                expected_revision=report.assessment_revision,
            )
            assert resolved.report.assessment_revision == report.assessment_revision + 1
            assert resolved.report.pending_review_count == 0
            assert resolved.report.status == "needs_review"
            assert resolved.report.flag_deduction == 0
            assert result.detail["n_source_reference_groups"] == 2
            assert result.detail["n_retained_reference_rows"] == 1
            assert "n_claim_support_candidate_pairs_total" not in result.detail
            metadata.assert_not_called()
            assert model.prompts == []
        else:
            assert structural.outcome == result.outcome == ResultOutcome.passed
            assert report.status == "ready" and report.flag_deduction == 0
            assert result.detail["n_claim_support_checked"] == 2
            assert len(result.detail["claim_support_occurrences"]) == 2
            assert metadata.await_count == 2 and model.prompts
        run_id, result_id, owner_id = run.id, result.id, owner.id
        detail = result.detail.copy()
    async with session_factory() as session:
        saved = await session.get(CheckResult, result_id)
        assert saved.detail == detail
        report = await get_report(session, run_id, owner_id)
        assert report.pending_review_count == 0
        assert report.status == ("needs_review" if partial else "ready")
        if partial:
            assert [item.coverage_reason for item in report.integrity_check_status] == [
                "reference_source_projection_unbound"
            ]
        audit = await session.scalar(
            select(AuditLog)
            .where(
                AuditLog.check_run_id == run_id,
                AuditLog.event_type == "citation_integrity_check_computed",
            )
            .order_by(AuditLog.id.desc())
            .limit(1)
        )
        assert audit.payload["reference_inventory_status"] == (
            "not_captured" if partial else "bound"
        )
        assert audit.payload["n_flags"] == 0
    assert path.read_bytes() == original
