"""Source-bound native math, safe prose and continuity controls (BUG-457)."""

import hashlib
from types import SimpleNamespace

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.checks.agreement.extract import extract_statements
from app.checks.forensics.service import _compute_forensics_findings
from app.checks.reuse.embed import _scoped_passages
from app.checks.semantic import _verify_quotes, build_semantic_batches
from app.config import Settings
from app.ingest.analytical_source import (
    analytical_source,
    bind_analytical_source,
    section_order_source_complete,
)
from app.ingest.docx import extract_document
from app.ingest.schemas import ExtractionResult
from app.models.enums import ResultOutcome


def _math(paragraph, text="x + 1 = 2", *, display=False):
    expression = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    content = OxmlElement("m:t")
    content.text = text
    run.append(content)
    expression.append(run)
    if display:
        wrapper = OxmlElement("m:oMathPara")
        wrapper.append(expression)
        paragraph._p.append(wrapper)
    else:
        paragraph._p.append(expression)
    return expression


def _extract(tmp_path, document):
    path = tmp_path / "source.docx"
    document.save(path)
    result = extract_document(str(path), Settings())
    assert result.docx_source_sha256 == hashlib.sha256(path.read_bytes()).hexdigest()
    # Validate the persisted raw-store form, not only in-memory model objects.
    return ExtractionResult.model_validate_json(result.model_dump_json())


def test_inline_math_never_becomes_a_joined_statistic_or_display_sentence(tmp_path):
    document = Document()
    document.add_paragraph("Independent readable evidence before.")
    mixed = document.add_paragraph("p = .")
    _math(mixed, "5")
    mixed.add_run("001")
    document.add_paragraph("t(28) = 2.45, p = .021")
    result = _extract(tmp_path, document)
    source = analytical_source(result)
    assert source.state == "complete"
    assert source.occurrence_count == 1
    assert [b.paragraph for b in result.blocks] == [0, 2]
    assert [[b.paragraph for b in scope] for scope in source.scopes] == [[0], [2]]
    container = source.inventory.containers[0]
    assert [(p.kind, p.text, p.occurrence) for p in container.pieces] == [
        ("text", "p = .", None),
        ("native_math", None, 0),
        ("text", "001", None),
    ]
    assert container.paragraph == 1
    assert "x + 1" not in result.model_dump_json()
    assert result.equations == []


def test_display_wrapper_multiple_and_fraction_are_expression_occurrences(tmp_path):
    document = Document()
    paragraph = document.add_paragraph()
    expression = _math(paragraph, display=True)
    fraction = OxmlElement("m:f")
    for tag in ("m:num", "m:den"):
        part = OxmlElement(tag)
        part.append(OxmlElement("m:r"))
        fraction.append(part)
    expression.append(fraction)
    _math(paragraph)
    result = _extract(tmp_path, document)
    source = analytical_source(result)
    assert source.occurrence_count == 2
    assert len(source.inventory.containers) == 1
    assert source.inventory.containers[0].paragraph == 0
    assert result.blocks == []
    assert result.text_chars == 0


def test_physical_merged_and_nested_cells_count_once_and_fence_only_affected_table(tmp_path):
    document = Document()
    table = document.add_table(rows=1, cols=2)
    cell = table.cell(0, 0).merge(table.cell(0, 1))
    _math(cell.paragraphs[0])
    nested = cell.add_table(rows=1, cols=1)
    _math(nested.cell(0, 0).paragraphs[0], display=True)
    safe = document.add_table(rows=2, cols=3)
    for row, values in zip(
        safe.rows, [("Group", "n", "M"), ("Control", "10", "3.33")], strict=True
    ):
        for target, value in zip(row.cells, values, strict=True):
            target.text = value
    result = _extract(tmp_path, document)
    source = analytical_source(result)
    assert source.occurrence_count == 2
    assert len(source.inventory.containers) == 2
    assert all(
        c.kind == "table_cell" and c.cell_path is not None for c in source.inventory.containers
    )
    assert [table.paragraph for table in source.tables] == [1]
    assert source.excluded_paragraphs == (0,)
    computation = _compute_forensics_findings(result)
    assert computation.n_descriptive == 2
    assert [flag.detail["kind"] for flag in computation.flag_drafts] == ["grim_inconsistent"]
    payload = result.model_dump(mode="json")
    payload["native_math_inventory"]["containers"][0]["cell_path"] = []
    assert analytical_source(ExtractionResult.model_validate(payload)).state == "unavailable"


def test_empty_fresh_inventory_is_zero_but_missing_is_unknown(tmp_path):
    document = Document()
    document.add_paragraph("A qualitative manuscript with readable source.")
    result = _extract(tmp_path, document)
    assert analytical_source(result).occurrence_count == 0
    assert analytical_source(result).scopes[0][0].text.startswith("A qualitative")
    legacy = result.model_copy(update={"native_math_inventory": None})
    unknown = analytical_source(legacy)
    assert unknown.state == "not_recorded"
    assert unknown.occurrence_count is None
    assert unknown.scopes == ()


def test_corrupt_or_wrong_source_inventory_cannot_certify_analytical_text(tmp_path):
    document = Document()
    _math(document.add_paragraph("ordinary prose"))
    result = _extract(tmp_path, document)
    payload = result.model_dump(mode="json")
    payload["native_math_inventory"]["occurrences"].append(
        payload["native_math_inventory"]["occurrences"][0]
    )
    corrupt = ExtractionResult.model_validate(payload)
    assert analytical_source(corrupt).state == "unavailable"
    assert analytical_source(corrupt).occurrence_count is None
    payload = result.model_dump(mode="json")
    payload["native_math_inventory"].update(occurrences=[], containers=[], display_cells=[])
    assert analytical_source(ExtractionResult.model_validate(payload)).state == "unavailable"
    assert analytical_source(result, expected_source_sha256="0" * 64).state == "unavailable"


def test_corrupt_body_ordinal_or_negative_xml_child_cannot_reconnect_scopes(tmp_path):
    document = Document()
    document.add_paragraph("The outcome was")
    _math(document.add_paragraph())
    document.add_paragraph("statistically significant.")
    result = _extract(tmp_path, document)
    for field in ("ordinal", "path"):
        payload = result.model_dump(mode="json")
        if field == "ordinal":
            payload["native_math_inventory"]["containers"][0]["paragraph"] = 9999
        else:
            payload["native_math_inventory"]["occurrences"][0]["xml_path"][-1] = -1
        invalid = analytical_source(ExtractionResult.model_validate(payload))
        assert invalid.state == "unavailable"
        assert invalid.scopes == ()
        assert invalid.occurrence_count is None


def test_textbox_math_preserves_outer_prose_without_certifying_continuity(tmp_path):
    document = Document()
    outer = document.add_paragraph("Ordinary preceding prose. " * 20)
    picture = OxmlElement("w:pict")
    content = OxmlElement("w:txbxContent")
    inner = OxmlElement("w:p")
    content.append(inner)
    picture.append(content)
    outer.add_run()._r.append(picture)
    from docx.text.paragraph import Paragraph

    _math(Paragraph(inner, None))
    outer.add_run(" Ordinary following prose. " * 20)
    result = _extract(tmp_path, document)
    source = analytical_source(result)
    assert source.occurrence_count == 1
    assert source.inventory.containers[0].paragraph == 0
    fragments = [
        piece.text for piece in source.inventory.containers[0].pieces if piece.kind == "text"
    ]
    assert any("preceding prose" in text for text in fragments)
    assert any("following prose" in text for text in fragments)
    assert result.text_chars > 900
    assert result.blocks == []


def test_inventory_scan_includes_body_wrappers_without_inventing_body_ordinal(tmp_path):
    document = Document()
    paragraph = document.add_paragraph("wrapped")
    _math(paragraph)
    body = document.element.body
    body.remove(paragraph._p)
    wrapper = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    content.append(paragraph._p)
    wrapper.append(content)
    body.insert(0, wrapper)
    document.add_paragraph("Independent paragraph.")
    result = _extract(tmp_path, document)
    source = analytical_source(result)
    assert source.state == "complete"
    assert source.occurrence_count == 1
    assert source.unlocated_occurrences == 1
    assert source.inventory.containers[0].paragraph is None
    assert source.scopes == ()
    assert source.inventory.occurrences[0].xml_path[0] == 0
    assert paragraph._p.tag == qn("w:p")


def test_real_neighbor_stat_and_wrong_p_flag_survive_native_omission(tmp_path):
    document = Document()
    document.add_paragraph("t(28) = 2.45, p = .500")
    _math(document.add_paragraph())
    result = _extract(tmp_path, document)
    computation = _compute_forensics_findings(result)
    assert computation.outcome == ResultOutcome.escalated
    assert computation.n_inferential == 1
    assert len(computation.flag_drafts) == 1
    assert computation.flag_drafts[0].detail["kind"] == "p_value_decision_error"
    assert computation.image_coverage["n_equations_unassessed"] == 0
    assert computation.image_coverage["n_native_equations_unassessed"] == 1
    assert computation.image_coverage["analysis_coverage"] == "partial"


def test_math_only_is_unsupported_not_no_statistics_or_passed(tmp_path):
    document = Document()
    _math(document.add_paragraph())
    computation = _compute_forensics_findings(_extract(tmp_path, document))
    assert computation.outcome == ResultOutcome.escalated
    assert computation.image_coverage["analysis_coverage"] == "unsupported"
    assert computation.image_coverage["n_native_equations_unassessed"] == 1


def test_gap_quote_rejected_without_losing_independent_quote(tmp_path):
    document = Document()
    document.add_paragraph("The outcome was")
    _math(document.add_paragraph())
    document.add_paragraph("statistically significant.")
    result = _extract(tmp_path, document)
    criterion = SimpleNamespace(text="Discuss the research outcome.", evidence=None)
    batches, missing = build_semantic_batches([criterion], result, settings=Settings())
    assert missing == []
    batch = batches[0][0]
    assert (
        _verify_quotes(
            ["The outcome was statistically significant."],
            batch.blocks,
            "paragraph",
            scopes=batch.source_scopes,
        )
        is None
    )
    assert _verify_quotes(
        ["statistically significant."],
        batch.blocks,
        "paragraph",
        scopes=batch.source_scopes,
    ) == ["paragraph 2"]
    assert '"contiguous_source_scopes"' in batch.context_text
    assert "cannot establish absence" in batch.context_text


def test_f4_sentence_cannot_continue_through_math_container(tmp_path):
    document = Document()
    document.add_paragraph("This study aims to")
    _math(document.add_paragraph())
    document.add_paragraph("automate enrollment and improve registration.")
    document.add_paragraph("Results showed that attendance improved.")
    result = _extract(tmp_path, document)
    statements = extract_statements(result, settings=Settings())
    assert all("automate enrollment" not in statement.text for statement in statements.intents)
    assert any("attendance improved" in statement.text for statement in statements.outcomes)


def test_f7_passage_and_context_never_cross_native_math_gap(tmp_path):
    document = Document()
    document.add_paragraph("Before the native equation.")
    _math(document.add_paragraph())
    document.add_paragraph("After the native equation.")
    result = _extract(tmp_path, document)
    drafts = _scoped_passages(
        analytical_source(result),
        result.blocks,
        chunk_words=100,
        ref_ids=set(),
        quote_ids=set(),
        context_words=100,
    )
    assert [draft["text"] for draft in drafts] == [
        "Before the native equation.",
        "After the native equation.",
    ]
    assert [draft["context_text"] for draft in drafts] == [draft["text"] for draft in drafts]
    assert [draft["paragraph"] for draft in drafts] == [0, 2]


def test_legacy_f6_count_unknown_and_known_empty_control_not_applicable(tmp_path):
    document = Document()
    document.add_paragraph("A qualitative case study describes interviews.")
    result = _extract(tmp_path, document)
    known = _compute_forensics_findings(result)
    assert known.outcome == ResultOutcome.not_applicable
    assert known.image_coverage["n_native_equations_unassessed"] == 0
    unknown = _compute_forensics_findings(result.model_copy(update={"native_math_inventory": None}))
    assert unknown.outcome == ResultOutcome.escalated
    assert unknown.image_coverage["n_native_equations_unassessed"] is None
    assert unknown.image_coverage["native_math_inventory_state"] == "not_recorded"


def test_expected_manuscript_binding_cannot_be_replaced_by_internal_extraction_hash(tmp_path):
    document = Document()
    document.add_paragraph("t(28) = 2.45, p = .021")
    result = _extract(tmp_path, document)
    saved = result.model_dump_json()
    bound = bind_analytical_source(result, "0" * 64)
    assert analytical_source(bound).state == "unavailable"
    assert _compute_forensics_findings(bound).n_inferential == 0
    assert result.model_dump_json() == saved
    assert "analytical_source_sha256" not in bound.model_dump_json()


def test_selected_archive_passage_and_context_must_share_one_safe_scope(tmp_path):
    from app.checks.reuse.source import selected_passage_supported

    document = Document()
    document.add_paragraph("The outcome was")
    _math(document.add_paragraph())
    document.add_paragraph("statistically significant.")
    result = _extract(tmp_path, document)
    source = analytical_source(result)
    assert selected_passage_supported(
        source,
        text="statistically significant.",
        context="statistically significant.",
        paragraph=2,
    )
    assert not selected_passage_supported(
        source,
        text="The outcome was statistically significant.",
        context="The outcome was statistically significant.",
        paragraph=0,
    )
    assert not selected_passage_supported(
        source,
        text="statistically significant.",
        context="The outcome was statistically significant.",
        paragraph=2,
    )
    assert not selected_passage_supported(
        analytical_source(result.model_copy(update={"native_math_inventory": None})),
        text="statistically significant.",
        context="statistically significant.",
        paragraph=2,
    )


@pytest.mark.parametrize(
    "paragraphs,text,context,paragraph,expected",
    [
        (
            ["Repeated passage", "Repeated passage"],
            "Repeated passage",
            "Repeated passage",
            0,
            False,
        ),
        (["aaaaa"], "aaa", "aaaaa", 0, True),
        (
            ["prefix unique", "following context"],
            "unique following",
            "unique following context",
            0,
            True,
        ),
        (["prefix", "unique passage"], "unique passage", "prefix unique passage", 1, True),
        (["prefix", "unique passage"], "unique passage", "prefix unique passage", 0, False),
    ],
)
def test_passage_locations_preserve_ambiguity_overlaps_and_paragraph_boundaries(
    tmp_path, paragraphs, text, context, paragraph, expected
):
    from app.checks.reuse.source import selected_passage_supported

    document = Document()
    for value in paragraphs:
        document.add_paragraph(value)
    source = analytical_source(_extract(tmp_path, document))
    assert (
        selected_passage_supported(source, text=text, context=context, paragraph=paragraph)
        is expected
    )


def test_section_order_cannot_pass_shortened_heading_tree_but_keeps_independent_failure(tmp_path):
    from app.checks.rules import RuleContext
    from app.checks.rules.sections import _run_section_order

    for numbers, omitted, expected in [
        ([1, 4, 2], None, ResultOutcome.failed),
        ([1, 4, 2], 4, ResultOutcome.unverifiable),
        ([1, 4, 3, 2], 4, ResultOutcome.failed),
        ([1, 2], None, ResultOutcome.passed),
        ([1, 4], 4, ResultOutcome.unverifiable),
    ]:
        document = Document()
        for number in numbers:
            heading = document.add_heading(f"Chapter {number} Research", level=1)
            if number == omitted:
                _math(heading)
        result = _extract(tmp_path, document)
        context = RuleContext(
            manuscript_id=1,
            anchor_kind="paragraph",
            page_count=0,
            section_tree=result.section_tree,
            geometry=[],
            tables=[],
            vision_status="none",
            has_images=False,
            citations=[],
            margin_tolerance_pts=0,
            citation_style_min_ratio=0,
            table_caption_min_ratio=0,
            section_order_source_complete=section_order_source_complete(result),
            native_math_source_detail=analytical_source(result).detail(),
        )
        outcome = _run_section_order(SimpleNamespace(text="Section order", evidence=None), context)
        assert outcome.outcome == expected
        if expected == ResultOutcome.unverifiable:
            assert "native Word equations" in outcome.detail["reason"]


def test_authorized_source_dto_keeps_safe_cell_fragments_without_private_xml(tmp_path, monkeypatch):
    from app.report import service

    document = Document()
    table = document.add_table(rows=1, cols=2)
    cell = table.cell(0, 0).merge(table.cell(0, 1))
    paragraph = cell.paragraphs[0]
    paragraph.add_run("p = .")
    _math(paragraph, "PRIVATE FORMULA")
    paragraph.add_run("001")
    result = _extract(tmp_path, document)
    monkeypatch.setattr(service, "load_raw_store", lambda settings, manuscript_id: result)
    manuscript = SimpleNamespace(
        id=1,
        file_ref=str(tmp_path / "source.docx"),
        purged_at=None,
        content_hash=result.docx_source_sha256,
        section_tree=result.section_tree.model_dump(mode="json"),
    )
    public = service.manuscript_paragraphs_for(manuscript, Settings())
    assert public.native_math.occurrence_count == 1
    assert public.tables[0].rows == [["", ""]]
    assert public.tables[0].native_cells[0].display_coordinates == [(0, 0), (0, 1)]
    assert [piece.kind for piece in public.tables[0].native_cells[0].pieces] == [
        "text",
        "native_math",
        "text",
    ]
    payload = public.model_dump_json()
    for private in ("PRIVATE FORMULA", "xml_path", "xml_sha256", "cell_path", "source_sha256"):
        assert private not in payload
    damaged = result.model_dump(mode="json")
    damaged["native_math_inventory"]["display_cells"][0]["display_coordinates"] = [[0, 5]]
    assert analytical_source(ExtractionResult.model_validate(damaged)).state == "unavailable"


def test_reference_source_scope_keeps_native_math_outside_the_bibliography_local(tmp_path):
    from app.ingest.patterns import load_patterns
    from app.ingest.reference_inventory import reference_source_scope_issues
    from app.ingest.references import select_references

    for inside in (False, True):
        document = Document()
        document.add_heading("Introduction", level=1)
        document.add_paragraph("Independent manuscript body.")
        if not inside:
            _math(document.add_paragraph())
        document.add_heading("References", level=1)
        document.add_paragraph("Smith, J. (2024). A named research source. Research Journal.")
        if inside:
            _math(document.add_paragraph())
        document.add_heading("Appendix", level=1)
        document.add_paragraph("Independent appendix content.")
        result = _extract(tmp_path, document)
        selected = select_references(result, load_patterns(Settings().ingest_patterns_file))
        issues = reference_source_scope_issues(result, selected)
        assert issues == (("reference_native_math_unassessed",) if inside else ())


def test_native_and_image_source_limitations_retain_their_separate_units(tmp_path):
    from io import BytesIO

    from PIL import Image

    document = Document()
    document.add_paragraph("t(28) = 2.45, p = .021")
    _math(document.add_paragraph())
    image = BytesIO()
    Image.new("RGB", (32, 32), "white").save(image, format="PNG")
    document.add_picture(BytesIO(image.getvalue()))
    computation = _compute_forensics_findings(_extract(tmp_path, document))
    coverage = computation.image_coverage
    assert computation.outcome == ResultOutcome.escalated
    assert computation.n_inferential == 1
    assert coverage["n_native_equations_unassessed"] == 1
    assert coverage["n_equations_unassessed"] == 0
    assert coverage["n_images_unassessed"] == 1
    assert coverage["coverage_reason"] == "image_content_unassessed"
    assert coverage["analysis_coverage"] == "partial"


@pytest.mark.parametrize("numbers", [[1, 2], [4, 2]])
@pytest.mark.parametrize("matched", [True, False])
async def test_actual_context_and_executor_do_not_grade_another_source(
    tmp_path, monkeypatch, numbers, matched
):
    from app.checks.router import route_criterion
    from app.checks.rules import context as context_module
    from app.checks.structural import run_structural_check

    document = Document()
    for number in numbers:
        document.add_heading(f"Chapter {number} Research", level=1)
    extraction = _extract(tmp_path, document)

    async def load(*args):
        return extraction

    class CaptureSession:
        def __init__(self):
            self.results = []

        async def execute(self, statement):
            return SimpleNamespace(scalars=lambda: SimpleNamespace(all=lambda: []))

        def add(self, result):
            self.results.append(result)

        async def commit(self):
            pass

    monkeypatch.setattr(context_module, "load_raw_store_async", load)
    session = CaptureSession()
    manuscript = SimpleNamespace(
        id=1, content_hash=extraction.docx_source_sha256 if matched else "0" * 64
    )
    ctx = await context_module.build_rule_context(session, manuscript, Settings())
    criterion = SimpleNamespace(text="Section order", evidence=None)
    decision = route_criterion(criterion, criterion_id=1, raw_type="structural")
    result = await run_structural_check(session, 1, criterion, 1, decision, ctx)
    expected = (
        (ResultOutcome.passed if numbers == [1, 2] else ResultOutcome.failed)
        if matched
        else ResultOutcome.unverifiable
    )
    assert result.outcome == expected
    assert result.score == ({ResultOutcome.passed: 100.0, ResultOutcome.failed: 0.0}.get(expected))
    assert session.results == [result]


@pytest.mark.parametrize("native", [False, True])
@pytest.mark.parametrize("ambiguous", [False, True])
async def test_f5_retains_supported_metadata_without_absence_from_an_omitted_citation(
    tmp_path, monkeypatch, native, ambiguous
):
    from unittest.mock import AsyncMock

    from app.checks.citations import verify
    from app.external.schemas import VerificationResult
    from app.ingest.patterns import load_patterns
    from app.ingest.references import extract_references
    from app.models.citation import Citation
    from app.models.run import Flag
    from tests.test_stabilization_integrity_controls import MemorySession

    document = Document()
    document.add_heading("Introduction", 1)
    document.add_paragraph("Prior work (Smith, 2024) found improvements.")
    if ambiguous:
        document.add_paragraph("Independent work (Adams, 2022) found another improvement.")
    omitted = document.add_paragraph("Related work (Jones, 2023) found another result.")
    if native:
        _math(omitted)
    document.add_heading("References", 1)
    document.add_paragraph("Smith, J. (2024). First source. Journal of Examples.")
    document.add_paragraph("Jones, A. (2023). Second source. Journal of Examples.")
    if ambiguous:
        document.add_paragraph("Smith, J. (2024). Another same-year source. Journal of Examples.")
        document.add_paragraph("Adams, J. (2022). Independent source. Journal of Examples.")
    extraction = _extract(tmp_path, document)
    citations = [
        Citation(manuscript_id=1, **vars(draft))
        for draft in extract_references(extraction, load_patterns())
    ]
    provider = AsyncMock(
        return_value=verify.CitationVerdict(
            verify.VerdictKind.retracted, VerificationResult(found=True, provider="control")
        )
    )
    monkeypatch.setattr(verify, "verify_citation", provider)
    monkeypatch.setattr(verify, "write_audit_event", AsyncMock())
    session = MemorySession()
    result = await verify.run_citation_integrity_check(
        session, None, 1, citations, extraction, load_patterns(), Settings(), None
    )
    flags = [row for row in session.rows if isinstance(row, Flag)]
    assert provider.await_count == (4 if ambiguous else 2)
    assert len([row for row in flags if row.detail["kind"] == "retracted_source"]) == (
        4 if ambiguous else 2
    )
    assert not any(row.detail["kind"] == "uncited_reference" for row in flags)
    # Ambiguous Smith is not a confirmed link; the independent Adams replaces
    # that positive control while both Smith reference indexes stay unassessed.
    assert result.detail["n_linked"] == 2 - int(native)
    assert result.detail["n_linkage_unassessed"] == (2 if ambiguous else 0) + int(native)
    if native:
        assert result.outcome == ResultOutcome.escalated
        assert result.detail["analysis_coverage"] == "partial"
        assert result.detail["coverage_reason"] == "native_math_citation_incomplete"
        from app.report.schemas import COVERAGE_REASONS

        assert result.detail["coverage_reason"] in COVERAGE_REASONS


@pytest.mark.parametrize("reported, expected_flags", [(".011", 0), (".500", 1)])
def test_f6_excluded_tail_declaration_cannot_create_a_false_mismatch(
    tmp_path, reported, expected_flags
):
    document = Document()
    document.add_heading("Methods", 1)
    methods = document.add_paragraph("All reported significance tests are one-tailed tests.")
    _math(methods, "alpha = .05")
    document.add_heading("Results", 1)
    document.add_paragraph(f"t(28) = 2.45, p = {reported}")
    computation = _compute_forensics_findings(_extract(tmp_path, document))
    assert computation.outcome == ResultOutcome.escalated
    assert computation.n_inferential == 1
    assert len(computation.flag_drafts) == expected_flags
    if expected_flags:
        draft = computation.flag_drafts[0]
        assert draft.detail["tail_context_complete"] is False
        assert draft.detail["kind"] == "p_value_decision_error"
        assert "one-tailed recalculation" in draft.detail["reason"]


@pytest.mark.parametrize("suffix", [".docx", ".pdf"])
@pytest.mark.parametrize("legacy", [False, True])
async def test_selected_docx_source_uses_retained_capability_not_filename(
    tmp_path, monkeypatch, suffix, legacy
):
    from app.checks.reuse.source import (
        coarse_candidate_supported,
        selected_native_source,
        selected_passage_supported,
    )

    document = Document()
    document.add_paragraph("Independently readable source.")
    _math(document.add_paragraph())
    extraction = _extract(tmp_path, document)
    if legacy:
        extraction.native_math_inventory = None
    row = SimpleNamespace(
        id=1,
        purged_at=None,
        file_ref=f"renamed{suffix}",
        content_hash=extraction.docx_source_sha256,
    )

    async def read(*args):
        return extraction

    async def get(*args):
        return row

    monkeypatch.setattr("app.checks.reuse.source.read_source", read)
    cache = {}
    source = await selected_native_source(SimpleNamespace(get=get), 1, Settings(), cache)
    assert (
        selected_passage_supported(
            source,
            text="Independently readable source.",
            context="Independently readable source.",
            paragraph=0,
        )
        is not legacy
    )
    assert not await coarse_candidate_supported(SimpleNamespace(get=get), 1, Settings(), cache)


def test_incomplete_tail_context_cannot_claim_an_ambiguous_significance_change(tmp_path):
    document = Document()
    omitted = document.add_paragraph("All tests are one-tailed.")
    _math(omitted)
    document.add_paragraph("t(28) = 2.00, p = .040")
    result = _compute_forensics_findings(_extract(tmp_path, document))
    assert len(result.flag_drafts) == 1
    assert result.flag_drafts[0].detail["kind"] == "p_value_mismatch"
    assert result.flag_drafts[0].severity.value == "med"
    assert "changes whether" not in result.flag_drafts[0].detail["reason"]


async def test_selected_pdf_with_word_filename_retains_its_native_capability(tmp_path, monkeypatch):
    from app.checks.reuse.source import (
        coarse_candidate_supported,
        selected_native_source,
        selected_passage_supported,
    )
    from app.ingest import pdf
    from tests.test_ingest_pdf import PdfBuilder

    path = tmp_path / "renamed.docx"
    builder = PdfBuilder()
    builder.new_page().line("Independent PDF source.")
    builder.save(path)
    extraction = pdf.extract_document(str(path), Settings())
    assert extraction.anchor_kind == "page"
    row = SimpleNamespace(
        id=1,
        purged_at=None,
        file_ref=str(path),
        content_hash=hashlib.sha256(path.read_bytes()).hexdigest(),
    )

    async def read(*args):
        return extraction

    async def get(*args):
        return row

    monkeypatch.setattr("app.checks.reuse.source.read_source", read)
    cache = {}
    source = await selected_native_source(SimpleNamespace(get=get), 1, Settings(), cache)
    assert selected_passage_supported(
        source, text="Independent PDF source.", context="Independent PDF source.", paragraph=None
    )
    assert await coarse_candidate_supported(SimpleNamespace(get=get), 1, Settings(), cache)
