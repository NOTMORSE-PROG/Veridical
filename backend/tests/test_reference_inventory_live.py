"""Fresh F3 judgments use actual DOCX ingestion, stored rows and PostgreSQL."""

import hashlib
import os
from decimal import Decimal

import pytest
from docx import Document
from sqlalchemy import delete, select

from app.checks.router import route_criterion
from app.checks.rules.context import build_rule_context
from app.checks.structural import run_structural_check
from app.config import Settings
from app.errors import ConflictError, NotFoundError
from app.ingest.service import ingest_manuscript
from app.models.audit import AuditLog
from app.models.citation import Citation
from app.models.enums import CheckKind, CheckRunStatus, ResultOutcome
from app.models.instructor import Instructor
from app.models.manuscript import Manuscript
from app.models.rubric import Criterion, Rubric
from app.models.run import CheckResult, CheckRun
from app.report.service import (
    aggregate_and_score,
    decide_report,
    get_report,
    list_escalated_for_run,
    resolve_escalation_for_run,
)
from tests.test_reference_inventory import REF_A, REF_B
from tests.test_report_get_live import _clean_tables as _clean_tables
from tests.test_report_get_live import scratch_url as scratch_url
from tests.test_report_get_live import session_factory as session_factory

pytestmark = pytest.mark.skipif(
    "DATABASE_URL" not in os.environ, reason="requires private Postgres"
)


@pytest.mark.parametrize(
    "case,expected",
    [
        ("complete", ResultOutcome.passed),
        ("missing_rows", ResultOutcome.unverifiable),
        ("partial_rows", ResultOutcome.unverifiable),
        ("table", ResultOutcome.unverifiable),
        ("known_zero", ResultOutcome.failed),
    ],
)
async def test_actual_ingestion_and_persisted_reference_criterion(
    session_factory, tmp_path, case, expected
):
    document = Document()
    document.add_paragraph("A synthetic study.")
    document.add_heading("References", 1)
    if case == "table":
        document.add_table(rows=1, cols=1).cell(0, 0).text = REF_A
    elif case != "known_zero":
        document.add_paragraph(REF_A)
        if case == "partial_rows":
            document.add_paragraph(REF_B)
    document.add_heading("Appendix", 1)
    path = tmp_path / "synthetic.docx"
    document.save(path)
    original = path.read_bytes()
    settings = Settings(_env_file=None, data_dir=tmp_path / "data")
    async with session_factory() as session:
        owner = Instructor(email="reference-control@example.test", display_name="Reference control")
        session.add(owner)
        await session.commit()
        manuscript = Manuscript(
            instructor_id=owner.id,
            group_label="Synthetic",
            file_ref=str(path),
            original_filename=path.name,
            content_hash=hashlib.sha256(original).hexdigest(),
        )
        rubric = Rubric(
            instructor_id=owner.id, title="Synthetic bound", source_file="synthetic.docx"
        )
        session.add_all([manuscript, rubric])
        await session.commit()
        await ingest_manuscript(session, manuscript, path, settings, read_images=False)
        criterion = Criterion(
            rubric_id=rubric.id,
            type="structural",
            text="At least 1 reference",
            evidence=None,
            weight=Decimal("1"),
            position=0,
        )
        session.add(criterion)
        await session.commit()
        if case in {"missing_rows", "partial_rows"}:
            condition = Citation.manuscript_id == manuscript.id
            if case == "partial_rows":
                condition = condition & (Citation.order_index == 1)
            await session.execute(delete(Citation).where(condition))
            await session.commit()
        run = CheckRun(manuscript_id=manuscript.id, rubric_id=rubric.id, status=CheckRunStatus.done)
        session.add(run)
        await session.commit()
        ctx = await build_rule_context(session, manuscript, settings)
        decision = route_criterion(criterion, criterion_id=criterion.id, raw_type="structural")
        result = await run_structural_check(session, run.id, criterion, criterion.id, decision, ctx)
        result_id, run_id, owner_id = result.id, run.id, owner.id
        assert result.outcome == expected
        assert result.score == (
            100
            if expected == ResultOutcome.passed
            else 0
            if expected == ResultOutcome.failed
            else None
        )
        if expected == ResultOutcome.unverifiable:
            assert "actual_count" not in result.detail
        elif expected == ResultOutcome.failed:
            assert result.detail["actual_count"] == 0
        await aggregate_and_score(session, run.id, settings, commit=True)
    async with session_factory() as session:
        saved = await session.get(CheckResult, result_id)
        assert saved.outcome == expected
        report = await get_report(session, run_id, owner_id)
        assert len(report.results) == 1 and report.results[0].outcome == expected.value
        if expected == ResultOutcome.unverifiable:
            assert report.composite_score is None and report.status == "needs_review"
            assert report.pending_review_count == 1
            # Integrity limits remain a separate lane: no criterion, no grade.
            integrity = CheckResult(
                check_run_id=run_id,
                criterion_id=None,
                kind=CheckKind.citation_integrity,
                outcome=ResultOutcome.unverifiable,
                detail={"coverage_reason": "reference_list_unavailable"},
            )
            session.add(integrity)
            await session.commit()
            pending = await list_escalated_for_run(session, run_id, owner_id)
            assert len(pending) == 1 and pending[0].check_result_id == result_id
            assert pending[0].review_reason == "not_graded"
            assert pending[0].agreement is None and pending[0].votes == []
            with pytest.raises(ConflictError, match="still need your review"):
                await decide_report(
                    session,
                    run_id,
                    owner_id,
                    "approved",
                    "Manual control reason.",
                    expected_revision=report.assessment_revision,
                )
            with pytest.raises(NotFoundError):
                await resolve_escalation_for_run(
                    session,
                    run_id,
                    result_id,
                    owner_id + 1,
                    "mark_pass",
                    "The instructor inspected the original reference list manually.",
                    expected_revision=report.assessment_revision,
                )
            # Atomic rejection rolls the transaction back and expires ORM state.
            await session.refresh(saved)
            original_detail = saved.detail.copy()
            resolved = await resolve_escalation_for_run(
                session,
                run_id,
                result_id,
                owner_id,
                "mark_pass",
                "The instructor inspected the original reference list manually.",
                expected_revision=report.assessment_revision,
            )
            assert resolved.report.assessment_revision == report.assessment_revision + 1
            assert resolved.report.pending_review_count == 0
            assert resolved.report.status == "needs_review"  # separate F5 limit retained
            await session.refresh(saved)
            assert saved.outcome == ResultOutcome.passed and saved.score == 100
            for key, value in original_detail.items():
                assert saved.detail[key] == value
            assert saved.detail["resolution"]["type"] == "mark_pass"
            assert await list_escalated_for_run(session, run_id, owner_id) == []
            await session.refresh(integrity)
            assert integrity.outcome == ResultOutcome.unverifiable
            audit = list(
                await session.scalars(select(AuditLog).where(AuditLog.check_run_id == run_id))
            )
            assert any(row.event_type == "escalation_resolved" for row in audit)
        elif expected == ResultOutcome.failed:
            assert report.composite_score == 0
    assert path.read_bytes() == original
