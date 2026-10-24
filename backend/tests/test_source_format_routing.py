"""F1/F8: accepted source bytes retain their identity across filename changes."""

import hashlib
from datetime import UTC, datetime
from types import SimpleNamespace

import pytest
from docx import Document

from app.config import Settings
from app.errors import ConflictError, GoneError
from app.ingest.docx import extract_document
from app.ingest.service import _write_raw_store, raw_store_path
from app.report.service import (
    _manuscript_export_source_format,
    _open_manuscript_pdf_or_none,
    _open_manuscript_source,
    manuscript_file_path_for,
    manuscript_paragraphs_for,
)
from app.report.source_files import manuscript_source_file, open_source_pdf
from tests.test_ingest_fixtures import FIXTURE_DIR
from tests.test_native_math_source import _math


@pytest.fixture()
def source(tmp_path):
    settings = Settings(_env_file=None, data_dir=tmp_path, storage_backend="local")
    manuscript = SimpleNamespace(
        id=1,
        purged_at=None,
        file_ref="",
        content_hash=None,
        section_tree={"source": "none", "nodes": []},
    )
    return settings, manuscript


@pytest.mark.parametrize("suffix", [".pdf", ".docx", ".txt", ""])
def test_pdf_identity_and_exact_original_ignore_filename(source, suffix):
    settings, manuscript = source
    path = settings.data_dir / f"original{suffix}"
    payload = (FIXTURE_DIR / "native.pdf").read_bytes()
    path.write_bytes(payload)
    manuscript.file_ref = str(path)
    doc, fmt, reason = _open_manuscript_source(manuscript, settings)
    assert fmt == "pdf" and reason is None
    try:
        assert doc.is_pdf and doc.page_count > 0
        assert doc[0].get_text().strip()
    finally:
        doc.close()
    assert manuscript_file_path_for(manuscript, settings).read_bytes() == payload
    assert _manuscript_export_source_format(manuscript, settings) == "pdf"
    with pytest.raises(ConflictError, match="isn't a DOCX"):
        manuscript_paragraphs_for(manuscript, settings)


@pytest.mark.parametrize("suffix", [".docx", ".pdf", ".txt", ""])
def test_docx_native_table_and_saved_prose_survive_misleading_name(source, suffix):
    settings, manuscript = source
    path = settings.data_dir / f"original{suffix}"
    document = Document()
    document.add_paragraph("Retained ordinary evidence.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Visible table evidence."
    _math(table.cell(0, 1).paragraphs[0], "PRIVATE EQUATION")
    document.save(path)
    extraction = extract_document(str(path), settings)
    manuscript.file_ref = str(path)
    manuscript.content_hash = extraction.docx_source_sha256
    manuscript.section_tree = extraction.section_tree.model_dump(mode="json")
    raw = raw_store_path(settings, manuscript.id)
    _write_raw_store(extraction, raw)
    before = hashlib.sha256(raw.read_bytes()).hexdigest()
    doc, fmt, reason = _open_manuscript_source(manuscript, settings)
    assert doc is None and fmt == "docx" and reason is None
    assert _manuscript_export_source_format(manuscript, settings) == "docx"
    assert _open_manuscript_pdf_or_none(manuscript, settings) is None
    public = manuscript_paragraphs_for(manuscript, settings)
    assert public.paragraphs[0].text == "Retained ordinary evidence."
    assert public.tables[0].rows[0][0] == "Visible table evidence."
    assert public.native_math.occurrence_count == 1
    assert len(public.tables[0].native_cells) == 1
    assert "PRIVATE EQUATION" not in public.model_dump_json()
    with pytest.raises(ConflictError, match="isn't a PDF"):
        manuscript_file_path_for(manuscript, settings)
    assert hashlib.sha256(raw.read_bytes()).hexdigest() == before


@pytest.mark.parametrize("payload", [b"not a source", b"%PDF-1.4 broken", b"PK\x03\x04broken"])
def test_unreadable_sources_never_stream_as_pdf(source, payload):
    settings, manuscript = source
    path = settings.data_dir / "original.pdf"
    path.write_bytes(payload)
    manuscript.file_ref = str(path)
    with pytest.raises(ConflictError):
        manuscript_file_path_for(manuscript, settings)
    doc, _, reason = _open_manuscript_source(manuscript, settings)
    assert doc is None and reason


def test_purge_and_missing_original_do_not_serve_retained_docx_text(source):
    settings, manuscript = source
    manuscript.file_ref = str(settings.data_dir / "gone.docx")
    with pytest.raises(GoneError):
        manuscript_paragraphs_for(manuscript, settings)
    manuscript.purged_at = datetime.now(UTC)
    with pytest.raises(GoneError, match="purged"):
        manuscript_source_file(manuscript, settings)
    doc, fmt, reason = _open_manuscript_source(manuscript, settings)
    assert doc is None and fmt == "unknown" and "deleted" in reason


def test_successful_office_native_open_is_not_pdf_identity(source):
    settings, _ = source
    path = settings.data_dir / "office.pdf"
    Document().save(path)
    with pytest.raises(ConflictError, match="as a PDF"):
        open_source_pdf(path)


@pytest.mark.parametrize("payload", [b"%PDF-1.4 fake", b"%PDF-1.4 broken"])
def test_retained_pdf_conflict_releases_native_stream_before_source_deletion(source, payload):
    settings, _ = source
    path = settings.data_dir / "malformed.pdf"
    path.write_bytes(payload)
    with pytest.raises(ConflictError) as captured:
        open_source_pdf(path)
    assert captured.value.__context__ is None
    # Retain the error, as an awaited request task does, during deletion.
    path.unlink()
    assert not path.exists()
