"""BUG-469: structural execution failures remain unassessed and reviewable."""

import os
from decimal import Decimal
from uuid import uuid4

import pytest
from docx import Document
from sqlalchemy import select, text
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

import app.checks.rules as rules
import app.checks.structural as structural
import app.pipeline.machine as machine
from app.checks.rules import RuleOutcome, RuleSpec
from app.checks.rules.context import build_rule_context
from app.config import get_settings
from app.db import sqlalchemy_url
from app.errors import ConflictError
from app.ingest.service import ingest_upload
from app.llm.fake import FakeLLMClient
from app.models.enums import CheckKind, CheckRunStatus, ResultOutcome
from app.models.instructor import Instructor
from app.models.rubric import Criterion, Rubric
from app.models.run import CheckResult, CheckRun, ReadinessReport
from app.report.service import (
    decide_report,
    get_report,
    list_escalated_for_run,
    resolve_escalation_for_run,
)

pytestmark = pytest.mark.skipif("DATABASE_URL" not in os.environ, reason="needs private Postgres")
SCRATCH_DB = "veridical_structural_failure_f300_test"


@pytest.fixture(scope="module")
def scratch_url():
    import asyncio

    from alembic import command
    from tests.test_schema import _admin_execute, _alembic_config, _swap_db

    base = os.environ["DATABASE_URL"]
    asyncio.run(_admin_execute(base, f'DROP DATABASE IF EXISTS "{SCRATCH_DB}"'))
    asyncio.run(_admin_execute(base, f'CREATE DATABASE "{SCRATCH_DB}"'))
    url = _swap_db(base, SCRATCH_DB)
    command.upgrade(_alembic_config(url), "head")
    yield url
    asyncio.run(_admin_execute(base, f'DROP DATABASE IF EXISTS "{SCRATCH_DB}" WITH (FORCE)'))


@pytest.fixture
async def session_factory(scratch_url):
    engine = create_async_engine(sqlalchemy_url(scratch_url))
    yield async_sessionmaker(engine, expire_on_commit=False)
    await engine.dispose()


@pytest.fixture(autouse=True)
async def clean_tables(session_factory):
    async with session_factory() as session:
        await session.execute(text("TRUNCATE audit_log, instructor RESTART IDENTITY CASCADE"))
        await session.commit()


async def _run_case(
    session_factory, tmp_path, monkeypatch, mode, *, include_pass=True, error_first=False
):
    """Actual ingest/routing/structural/persistence; integrity is a settled control."""
    monkeypatch.setenv("DATA_DIR", str(tmp_path / "data"))
    get_settings.cache_clear()
    settings = get_settings()
    source = tmp_path / "verified.docx"
    document = Document()
    document.add_heading("Synthetic structural-failure control", 0)
    document.add_paragraph(
        "This document supplies real source bytes for a controlled rule failure."
    )
    document.save(source)

    async def chunks():
        yield source.read_bytes()

    def passing(_criterion, _context):
        return RuleOutcome(outcome=ResultOutcome.passed, anchor="document", detail={})

    def second(_criterion, _context):
        if mode == "na":
            return RuleOutcome(
                outcome=ResultOutcome.not_applicable,
                anchor="document",
                detail={"reason": "This conditional requirement does not apply to this control."},
            )
        raise ZeroDivisionError("private-sentinel C:/private/internal-rule-path")

    monkeypatch.setattr(rules, "_REGISTRY", dict(rules._snapshot_registry()))
    for rule_id, label, runner in (
        ("controlled_pass", "Controlled successful check", passing),
        ("controlled_failure", "Controlled second check", second),
    ):
        rules.register_rule(
            RuleSpec(rule_id, "Controlled rule", lambda c, s=label: c.text == s, runner)
        )
    if mode == "missing":
        original = structural.get_rule
        monkeypatch.setattr(
            structural,
            "get_rule",
            lambda key: None if key == "controlled_failure" else original(key),
        )

    async def settled_integrity(session, run, manuscript, settings, llm, cancellation_boundary):
        # These are explicit controlled integrity outcomes, not an F4-F7 evaluation.
        # A fully settled denominator prevents unrelated missing checks hiding Ready.
        controls = (
            (CheckKind.internal_agreement, ResultOutcome.not_applicable, {}),
            (CheckKind.citation_integrity, ResultOutcome.passed, {"analysis_coverage": "complete"}),
            (CheckKind.statistical_forensics, ResultOutcome.not_applicable, {}),
            (
                CheckKind.originality_reuse,
                ResultOutcome.passed,
                {
                    "archive_size_n": 1,
                    "corpus": "library",
                    "comparison_ran": True,
                    "n_flags": 0,
                },
            ),
        )
        session.add_all(
            CheckResult(
                check_run_id=run.id,
                criterion_id=None,
                kind=kind,
                outcome=outcome,
                score=None,
                detail=detail,
            )
            for kind, outcome, detail in controls
        )
        machine._record_stage(run, CheckRunStatus.integrity, status="done")

    monkeypatch.setattr(machine, "_run_integrity_stage", settled_integrity)
    async with session_factory() as session:
        instructor = Instructor(email=f"failure-{uuid4().hex}@test.local", display_name="Control")
        session.add(instructor)
        await session.commit()
        manuscript, *_ = await ingest_upload(
            session, chunks(), source.name, "Control", instructor_id=instructor.id
        )
        context = await build_rule_context(session, manuscript, settings)
        assert context.source_identity_verified
        rubric = Rubric(
            instructor_id=instructor.id,
            title="Controlled rules",
            source_file="controlled.docx",
            is_active=True,
        )
        session.add(rubric)
        await session.commit()
        labels = (["Controlled successful check"] if include_pass else []) + [
            "Controlled second check"
        ]
        if error_first:
            labels.reverse()
        criteria = [
            Criterion(
                rubric_id=rubric.id,
                type=(
                    "not_assessable"
                    if mode == "not_assessable" and label == "Controlled second check"
                    else "structural"
                ),
                text=label,
                weight=Decimal("1"),
                position=i,
            )
            for i, label in enumerate(labels)
        ]
        session.add_all(criteria)
        await session.commit()
        run = CheckRun(manuscript_id=manuscript.id, rubric_id=rubric.id)
        session.add(run)
        await session.commit()
        ids = (
            instructor.id,
            run.id,
            next(c.id for c in criteria if c.text == "Controlled second check"),
        )
        await machine.run_check_run(session, run, settings, FakeLLMClient())
        assert run.status == CheckRunStatus.done
    return ids


@pytest.mark.parametrize("mode", ["raise", "missing"])
async def test_mixed_rule_failure_is_incomplete_pending_and_cannot_be_approved(
    session_factory, tmp_path, monkeypatch, mode, caplog
):
    owner, run_id, criterion_id = await _run_case(session_factory, tmp_path, monkeypatch, mode)
    async with session_factory() as session:
        result = await session.scalar(
            select(CheckResult).where(
                CheckResult.check_run_id == run_id, CheckResult.criterion_id == criterion_id
            )
        )
        report = await get_report(session, run_id, owner)
        assert report.status == "needs_review"
        assert result.outcome == ResultOutcome.unverifiable
        assert result.score is None
        assert "private-sentinel" not in str(result.detail)
        assert "internal-rule-path" not in str(result.detail)
        assert result.detail["failure_kind"] == (
            "rule_execution_failed" if mode == "raise" else "rule_not_registered"
        )
        assert report.pending_review_count == 1
        assert report.reason
        pending = await list_escalated_for_run(session, run_id, owner)
        assert [item.check_result_id for item in pending] == [result.id]
        assert pending[0].ai_majority_verdict is None
    if mode == "raise":
        assert "private-sentinel" in caplog.text
    for decision in ("approved", "rejected"):
        async with session_factory() as session:
            with pytest.raises(ConflictError, match="still need"):
                await decide_report(
                    session,
                    run_id,
                    owner,
                    decision,
                    "I reviewed the available results.",
                    expected_revision=0,
                )
    async with session_factory() as session:
        row = await session.scalar(
            select(ReadinessReport).where(ReadinessReport.check_run_id == run_id)
        )
        assert row.decision is None and row.assessment_revision == 0


@pytest.mark.parametrize("mode", ["na", "not_assessable"])
async def test_genuine_exclusions_remain_settled(session_factory, tmp_path, monkeypatch, mode):
    owner, run_id, criterion_id = await _run_case(session_factory, tmp_path, monkeypatch, mode)
    async with session_factory() as session:
        result = await session.scalar(
            select(CheckResult).where(
                CheckResult.check_run_id == run_id, CheckResult.criterion_id == criterion_id
            )
        )
        assert result.outcome == ResultOutcome.not_applicable and result.score is None
        report = await get_report(session, run_id, owner)
        assert report.status == "ready" and report.pending_review_count == 0
        approved = await decide_report(
            session, run_id, owner, "approved", None, expected_revision=0
        )
        assert approved.decision == "approved"


async def test_all_rule_errors_remain_reviewable_without_a_score(
    session_factory, tmp_path, monkeypatch
):
    owner, run_id, _ = await _run_case(
        session_factory, tmp_path, monkeypatch, "raise", include_pass=False
    )
    async with session_factory() as session:
        report = await get_report(session, run_id, owner)
        assert report.status == "needs_review" and report.composite_score is None
        assert report.pending_review_count == 1


async def test_rule_error_blocks_actual_approval_even_with_a_note(
    session_factory, tmp_path, monkeypatch
):
    owner, run_id, _ = await _run_case(session_factory, tmp_path, monkeypatch, "raise")
    async with session_factory() as session:
        with pytest.raises(ConflictError, match="still need"):
            await decide_report(
                session,
                run_id,
                owner,
                "approved",
                "I inspected the available results.",
                expected_revision=0,
            )


async def test_later_passing_criterion_runs_after_error(session_factory, tmp_path, monkeypatch):
    _, run_id, criterion_id = await _run_case(
        session_factory, tmp_path, monkeypatch, "raise", error_first=True
    )
    async with session_factory() as session:
        rows = (
            (
                await session.execute(
                    select(CheckResult).where(
                        CheckResult.check_run_id == run_id, CheckResult.criterion_id.isnot(None)
                    )
                )
            )
            .scalars()
            .all()
        )
        assert len(rows) == 2
        assert (
            next(r for r in rows if r.criterion_id == criterion_id).outcome
            == ResultOutcome.unverifiable
        )
        assert (
            next(r for r in rows if r.criterion_id != criterion_id).outcome == ResultOutcome.passed
        )


@pytest.mark.parametrize(
    "resolution,expected",
    [
        ("mark_pass", ResultOutcome.passed),
        ("mark_fail", ResultOutcome.failed),
        ("needs_document", ResultOutcome.not_applicable),
    ],
)
async def test_instructor_can_resolve_error_with_original_failure_preserved(
    session_factory, tmp_path, monkeypatch, resolution, expected
):
    owner, run_id, criterion_id = await _run_case(session_factory, tmp_path, monkeypatch, "raise")
    async with session_factory() as session:
        pending = await list_escalated_for_run(session, run_id, owner)
        assert len(pending) == 1
        result_id = pending[0].check_result_id
    for invalid_resolution, reason in (
        ("accept_majority", "I reviewed this missing model judgment."),
        (resolution, "x"),
    ):
        async with session_factory() as session:
            with pytest.raises(ConflictError):
                await resolve_escalation_for_run(
                    session,
                    run_id,
                    result_id,
                    owner,
                    invalid_resolution,
                    reason,
                    expected_revision=0,
                )
    reason = "I inspected the source and recorded my own judgment after the automatic check failed."
    async with session_factory() as session:
        resolved = await resolve_escalation_for_run(
            session, run_id, result_id, owner, resolution, reason, expected_revision=0
        )
        assert resolved.outcome == expected.value
        assert resolved.report.pending_review_count == 0
    async with session_factory() as session:
        stored = await session.get(CheckResult, result_id)
        assert stored.criterion_id == criterion_id
        assert stored.detail["original_result"]["outcome"] == "unverifiable"
        assert stored.detail["original_result"]["score"] is None
        assert stored.detail["failure_kind"] == "rule_execution_failed"
        assert stored.detail["resolution"]["reason"] == reason
        assert stored.score == (
            None
            if resolution == "needs_document"
            else Decimal("100")
            if resolution == "mark_pass"
            else Decimal("0")
        )
