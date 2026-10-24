"""Reject oversized directory allocation before constructing ZipFile."""

import struct
import zipfile

import pytest
from docx import Document

from app.config import Settings
from app.errors import FileMalformedError
from app.ingest.docx import _check_uncompressed_size
from app.ingest.service import detect_format
from app.ingest.zip_bounds import check_docx_directory_size


def _as_zip64(data):
    start = data.rfind(zipfile.stringEndArchive)
    end = list(
        struct.unpack(zipfile.structEndArchive, data[start : start + zipfile.sizeEndCentDir])
    )
    count, size, offset = end[4:7]
    record = struct.pack(
        zipfile.structEndArchive64,
        zipfile.stringEndArchive64,
        44,
        45,
        45,
        0,
        0,
        count,
        count,
        size,
        offset,
    )
    locator = struct.pack(
        zipfile.structEndArchive64Locator, zipfile.stringEndArchive64Locator, 0, start, 1
    )
    end[3:7] = [0xFFFF, 0xFFFF, 0xFFFFFFFF, 0xFFFFFFFF]
    return (
        data[:start]
        + record
        + locator
        + struct.pack(zipfile.structEndArchive, *end)
        + data[start + zipfile.sizeEndCentDir :]
    )


@pytest.mark.parametrize("variant", ["ordinary", "comment", "zip64"])
def test_normal_word_packages_keep_exact_directory_budget(tmp_path, variant):
    path = tmp_path / "source.pdf"
    document = Document()
    document.add_paragraph("Known ordinary Word source.")
    document.save(path)
    if variant == "comment":
        with zipfile.ZipFile(path, "a") as package:
            package.comment = b"A valid archive comment."
    if variant == "zip64":
        path.write_bytes(_as_zip64(path.read_bytes()))
    with path.open("rb") as stream:
        size = zipfile._EndRecData(stream)[zipfile._ECD_SIZE]
    settings = Settings(_env_file=None, docx_zip_directory_max_bytes=size)
    assert detect_format(path, settings) == ".docx"
    _check_uncompressed_size(str(path), settings)
    assert Document(path).paragraphs[0].text == "Known ordinary Word source."
    with pytest.raises(FileMalformedError):
        detect_format(path, settings.model_copy(update={"docx_zip_directory_max_bytes": size - 1}))


@pytest.mark.parametrize("entry", ["detector", "extractor"])
def test_many_empty_members_rejected_before_zipfile_allocation(tmp_path, monkeypatch, entry):
    path = tmp_path / "source.docx"
    Document().save(path)
    with zipfile.ZipFile(path, "a") as package:
        for index in range(1000):
            package.writestr(f"empty/{index}", b"")
    settings = Settings(_env_file=None, docx_zip_directory_max_bytes=4096)
    monkeypatch.setattr(
        zipfile, "ZipFile", lambda *args, **kwargs: pytest.fail("Allocated directory")
    )
    with pytest.raises(FileMalformedError):
        if entry == "detector":
            detect_format(path, settings)
        else:
            _check_uncompressed_size(str(path), settings)


def test_zip64_directory_size_cannot_hide_behind_small_regular_end_record(tmp_path, monkeypatch):
    path = tmp_path / "source.docx"
    Document().save(path)
    data = bytearray(_as_zip64(path.read_bytes()))
    start = data.index(zipfile.stringEndArchive64)
    record = list(
        struct.unpack(zipfile.structEndArchive64, data[start : start + zipfile.sizeEndCentDir64])
    )
    record[-2] = 1 << 32
    data[start : start + zipfile.sizeEndCentDir64] = struct.pack(
        zipfile.structEndArchive64, *record
    )
    path.write_bytes(data)
    monkeypatch.setattr(
        zipfile, "ZipFile", lambda *args, **kwargs: pytest.fail("Allocated directory")
    )
    with pytest.raises(FileMalformedError):
        detect_format(path, Settings(_env_file=None))


@pytest.mark.parametrize("payload", [b"PK\x03\x04", b"PK\x03\x04" + b"x" * 100])
def test_missing_end_metadata_fails_closed(tmp_path, payload):
    path = tmp_path / "source.docx"
    path.write_bytes(payload)
    with pytest.raises(FileMalformedError):
        check_docx_directory_size(path, Settings(_env_file=None))


def test_truncated_zip64_uses_real_file_seek_errors_and_fails_closed(tmp_path):
    path = tmp_path / "source.docx"
    locator = struct.pack(
        zipfile.structEndArchive64Locator, zipfile.stringEndArchive64Locator, 0, 0, 1
    )
    end = struct.pack(zipfile.structEndArchive, zipfile.stringEndArchive, 0, 0, 1, 1, 1, 0, 0)
    path.write_bytes(b"PK\x03\x04" + locator + end)
    assert path.stat().st_size == 46
    with pytest.raises(FileMalformedError):
        detect_format(path, Settings(_env_file=None))
