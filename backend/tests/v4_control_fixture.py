"""V-081: deterministic source for the separate F5/F6 mechanism control."""

import io
import json
from datetime import datetime
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import docx

FIXTURE_DIR = Path(__file__).parent / "fixtures" / "v4_source_control_v2"


def load_recipe() -> dict:
    return json.loads((FIXTURE_DIR / "recipe.json").read_text(encoding="utf-8"))


def load_manifest() -> dict:
    return json.loads((FIXTURE_DIR / "provenance.json").read_text(encoding="utf-8"))


def build_control_docx(destination: Path) -> Path:
    """Only input data drives the document; no observed result enters it.

    Normalize package times and properties so every test can identify the
    exact same bytes, rather than only a similar-looking rebuilt document.
    """
    recipe = load_recipe()
    document = docx.Document()
    document.add_paragraph(recipe["title"], style="Title")
    document.add_paragraph(recipe["purpose"])
    document.add_heading(recipe["literature_heading"], level=1)
    document.add_paragraph(recipe["literature_text"])
    document.add_heading(recipe["statistics_heading"], level=1)
    document.add_paragraph(recipe["statistics_scope"])
    document.add_paragraph(recipe["inferential_text"])
    document.add_paragraph(recipe["descriptive_text"])
    document.add_paragraph(recipe["table_caption"])
    rows = recipe["table_rows"]
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    for target, source in zip(table.rows, rows, strict=True):
        for cell, value in zip(target.cells, source, strict=True):
            cell.text = value
    document.add_heading(recipe["references_heading"], level=1)
    for reference in recipe["references"]:
        document.add_paragraph(reference)
    properties = document.core_properties
    properties.created = properties.modified = datetime.fromisoformat(recipe["package_created_at"])
    properties.author = properties.last_modified_by = ""
    properties.comments = ""
    properties.revision = 1
    content = io.BytesIO()
    document.save(content)
    content.seek(0)
    with ZipFile(content) as source, ZipFile(destination, "w", compression=ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            info = ZipInfo(name, date_time=tuple(recipe["package_zip_time"]))
            info.create_system = 0
            info.external_attr = 0o600 << 16
            info.compress_type = ZIP_DEFLATED
            target.writestr(info, source.read(name))
    return destination


def control_inventory(path: Path) -> dict:
    """Review the generated input without invoking any product analysis."""
    document = docx.Document(path)
    return {
        "paragraphs": [paragraph.text for paragraph in document.paragraphs],
        "tables": [
            [[cell.text for cell in row.cells] for row in table.rows] for table in document.tables
        ],
    }
