"""Adversarial stress test suite for DOCX pipeline & export parity.

Exercises:
1. Native DOCX tables with no captions evaluated against table_caption_min_ratio.
2. DOCX files with unusual file extensions, paths, and heuristics-based
   SectionTrees in manuscript_paragraphs_for.
3. Multi-page ReportLab PDF export with running headers/footers, zero layout
   clipping, and DOCX notices.
"""

import tempfile
from pathlib import Path

import docx as docx_lib
import fitz

from app.checks.rules import RuleContext, get_rule
from app.checks.rules.tables import RULE_ID as TABLE_FORMAT_RULE_ID
from app.config import get_settings
from app.ingest.docx import extract_document
from app.ingest.schemas import SectionTree, TableBlock
from app.ingest.service import _write_raw_store, raw_store_path
from app.models.enums import ResultOutcome
from app.models.manuscript import Manuscript
from app.report.export import build_report_pdf
from app.report.schemas import (
    CriterionResultOut,
    FlagExportDetailOut,
    FlagSummaryOut,
    PageEvidenceImageOut,
    ReportExportData,
    ReportOut,
)
from app.report.service import manuscript_paragraphs_for


class FakeCriterion:
    def __init__(self, text: str):
        self.text = text


def test_docx_native_tables_no_captions_stress():
    """Verify table_caption_min_ratio evaluates honestly without crashing or tautologies
    when native tables in DOCX have caption=None (fails) vs valid captions (passes).
    """
    spec = get_rule(TABLE_FORMAT_RULE_ID)
    criterion = FakeCriterion(text="Tables must be present and formatted with captions")

    # Multiple native DOCX tables with caption=None
    tables = [
        TableBlock(paragraph=2, rows=[["col1", "col2"], ["v1", "v2"]], caption=None),
        TableBlock(paragraph=5, rows=[["a", "b", "c"], ["1", "2", "3"]], caption=None),
        TableBlock(paragraph=8, rows=[["x", "y"], ["z", "w"]], caption=None),
    ]

    ctx = RuleContext(
        manuscript_id=1,
        anchor_kind="paragraph",
        page_count=0,
        section_tree=SectionTree(source="styles", nodes=[]),
        geometry=[],
        tables=tables,
        vision_status="none",
        has_images=False,
        citations=[],
        margin_tolerance_pts=5.0,
        citation_style_min_ratio=0.8,
        table_caption_min_ratio=0.8,
    )

    outcome = spec.run(criterion, ctx)
    assert outcome.outcome == ResultOutcome.failed
    assert outcome.detail["total_tables"] == 3
    assert outcome.detail["captioned"] == 0
    assert outcome.detail["ratio"] == 0.0

    # With captions, the same tables must pass
    captioned_tables = [
        TableBlock(paragraph=2, rows=[["col1", "col2"], ["v1", "v2"]], caption="Table 1: Metrics"),
        TableBlock(
            paragraph=5, rows=[["a", "b", "c"], ["1", "2", "3"]], caption="Table 2: Comparison"
        ),
        TableBlock(paragraph=8, rows=[["x", "y"], ["z", "w"]], caption="Table 3: Results"),
    ]
    passed_ctx = RuleContext(
        manuscript_id=1,
        anchor_kind="paragraph",
        page_count=0,
        section_tree=SectionTree(source="styles", nodes=[]),
        geometry=[],
        tables=captioned_tables,
        vision_status="none",
        has_images=False,
        citations=[],
        margin_tolerance_pts=5.0,
        citation_style_min_ratio=0.8,
        table_caption_min_ratio=0.8,
    )
    passed_outcome = spec.run(criterion, passed_ctx)
    assert passed_outcome.outcome == ResultOutcome.passed
    assert passed_outcome.detail["total_tables"] == 3
    assert passed_outcome.detail["captioned"] == 3


def test_docx_manuscript_paragraphs_heuristics_unusual_path():
    """Stress test: Realistic student DOCX (using heuristics / bold headings rather than styles)
    with unusual file extensions or paths in manuscript_paragraphs_for.
    """
    settings = get_settings()

    doc = docx_lib.Document()
    # Student document: plain paragraphs with bold headings, not Word styles
    p1 = doc.add_paragraph()
    r1 = p1.add_run("CHAPTER 1 INTRODUCTION")
    r1.bold = True
    doc.add_paragraph("This is body text of chapter 1.")

    temp_dir = Path(tempfile.mkdtemp())
    # Unusual file path / extension: storage blob without .docx suffix
    docx_path = temp_dir / "storage_blob_12345"
    doc.save(str(docx_path))

    ext = extract_document(docx_path, settings)
    manuscript_id = 88888
    store_p = raw_store_path(settings, manuscript_id)
    store_p.parent.mkdir(parents=True, exist_ok=True)
    _write_raw_store(ext, store_p)

    # Manuscript record with unusual file_ref and no extension in original_filename
    m = Manuscript(
        id=manuscript_id,
        original_filename="student_thesis_upload",
        file_ref=str(docx_path),
        section_tree=ext.section_tree.model_dump(),
        content_hash="hash_88888",
    )

    # SectionTree source is 'heuristics' (or 'none'), which should still be recognized as docx
    assert m.file_format == "docx", (
        f"Manuscript.file_format returned '{m.file_format}' for student DOCX with heuristics source"
    )

    paragraphs_out = manuscript_paragraphs_for(m, settings)
    assert len(paragraphs_out.paragraphs) >= 1


def test_docx_reportlab_multipage_export_stress():
    """Verify multi-page PDF report generation for DOCX manuscripts:
    ensure ReportLab exports clean multi-page documents with running headers/footers
    and zero layout clipping.
    """
    results = [
        CriterionResultOut(
            check_result_id=i + 1,
            criterion_id=i + 1,
            text=(
                f"Criterion {i + 1}: Required academic formatting standard "
                "for technical manuscripts."
            ),
            type="structural",
            weight=1.0,
            weight_importance="med",
            kind="rule",
            outcome="passed" if i % 3 != 0 else "failed",
            score=100.0 if i % 3 != 0 else 0.0,
            basis="Rule evaluation",
            anchor=f"paragraph {i * 4}",
            reasoning=f"Evaluated paragraph {i * 4} formatting rules.",
            reason="Complies with standard" if i % 3 != 0 else "Missing required elements.",
            evidence=[],
            resolution=None,
        )
        for i in range(25)
    ]

    flags = [
        FlagSummaryOut(
            id=101,
            check_kind="citations",
            severity="medium",
            criterion_text="Citation Integrity Verification",
            evidence_excerpt=(
                "(Smith et al., 2020) was cited in paragraph 12 but lacks bibliography entry."
            ),
            page_anchor="paragraph 12",
            overridden=False,
            affects_readiness=True,
        )
    ]

    flag_details = {
        101: FlagExportDetailOut(
            flag_id=101,
            finding_headline="Unlinked In-Text Citation",
            check_kind="citations",
            category="missing_citation",
            severity="medium",
            page_anchor="paragraph 12",
            affects_readiness=True,
            excerpt=(
                "(Smith et al., 2020) was cited in paragraph 12 but lacks bibliography entry."
            ),
            location_label="paragraph 12",
            reasoning=(
                "Citation integrity check scanned references and found no match for Smith (2020)."
            ),
            override_reason=None,
        )
    }

    page_images = {
        101: PageEvidenceImageOut(
            image_png=None,
            page=None,
            reason="DOCX layout - no page image",
            reason_code="docx_layout",
        )
    }

    report = ReportOut(
        check_run_id=42,
        manuscript_id=99,
        manuscript_group_label="Team Alpha BSIT Capstone 2026",
        manuscript_original_filename="capstone_final_submission.docx",
        rubric_title="TIP BSIT Capstone Defense Rubric 2026",
        status="conditionally_ready",
        composite_score=78.5,
        thresholds={"ready_min_score": 85.0, "not_ready_max_score": 60.0},
        reason="Criteria require review.",
        flag_deduction=5.0,
        unresolved_high_flag_count=0,
        llm_mode="real",
        results=results,
        decision=None,
        decided_at=None,
        decision_note=None,
        pending_review_count=1,
        rubric_is_current=True,
        previous_status=None,
        previous_composite_score=None,
    )

    export_data = ReportExportData(
        report=report,
        flags=flags,
        flag_details=flag_details,
        page_images=page_images,
        archive_size_n=15,
        source_format="docx",
    )

    pdf_bytes = build_report_pdf(export_data)
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    assert len(doc) >= 3, f"Expected at least 3 pages, got {len(doc)}"

    for i, page in enumerate(doc):
        text = page.get_text()
        rect = page.rect
        blocks = page.get_text("blocks")
        # Ensure no blocks are clipped outside page rect
        clipped = [
            b
            for b in blocks
            if b[0] < -0.1 or b[1] < -0.1 or b[2] > rect.width + 0.1 or b[3] > rect.height + 0.1
        ]
        assert len(clipped) == 0, f"Page {i + 1} has {len(clipped)} clipped text blocks"

        # Check footer on all pages
        assert f"Page {i + 1}" in text, f"Page {i + 1} missing footer page number"

        # Check running header on later pages
        if i > 0:
            assert "VERIDICAL - Readiness Report" in text, f"Page {i + 1} missing running header"

    full_text = "\n".join(p.get_text() for p in doc)
    assert "This manuscript was submitted as a DOCX" in full_text
