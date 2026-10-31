"""V-010: POST /rubrics — the HTTP surface, fake-LLM mode end-to-end
(DoD item 3). Needs a live Postgres (same convention as test_ingest_api.py).
"""

import os
from unittest.mock import AsyncMock

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import text

import app.auth.service as auth_service
from app.auth.security import hash_password
from app.config import get_settings
from tests.test_ingest_pdf import PdfBuilder

live = pytest.mark.skipif(
    "DATABASE_URL" not in os.environ,
    reason="integration: needs a live Postgres (CI service or local docker-compose)",
)
pytestmark = live

SCRATCH_DB = "veridical_rubricapitest"


@pytest.fixture(scope="module")
def api_scratch_url():
    import asyncio

    from alembic import command
    from tests.test_schema import _admin_execute, _alembic_config, _swap_db

    base = os.environ["DATABASE_URL"]
    asyncio.run(_admin_execute(base, f'DROP DATABASE IF EXISTS "{SCRATCH_DB}"'))
    asyncio.run(_admin_execute(base, f'CREATE DATABASE "{SCRATCH_DB}"'))
    url = _swap_db(base, SCRATCH_DB)
    command.upgrade(_alembic_config(url), "head")
    yield url
    asyncio.run(_admin_execute(base, f'DROP DATABASE IF EXISTS "{SCRATCH_DB}" WITH (FORCE)'))


@pytest.fixture()
def client(api_scratch_url, tmp_path, monkeypatch):
    import app.db as db
    import app.ratelimit as ratelimit

    monkeypatch.setenv("DATABASE_URL", api_scratch_url)
    monkeypatch.setenv("DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("VERIDICAL_FAKE_LLM", "1")
    get_settings.cache_clear()
    db._engine = None
    auth_service._rate_limiter = None
    ratelimit._limiters.clear()  # BUG-004: don't leak rubric-upload counts across tests
    from app.main import app

    with TestClient(app) as c:
        yield c
    db._engine = None


@pytest.fixture(autouse=True)
def _clean_tables(api_scratch_url):
    """Every test shares the same demo instructor (fixed email) in this
    module-scoped scratch DB — family-list assertions need a clean slate.
    Uses its OWN disposable engine, never `app.db.get_engine()` — that
    one is bound to whichever event loop first touches it, and mixing it
    with a separate `asyncio.run()` call corrupted it across loops on
    Windows before (same bug fixed in test_auth_router.py's `seeded`)."""
    import asyncio

    from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

    from app.db import sqlalchemy_url

    async def _truncate():
        engine = create_async_engine(sqlalchemy_url(api_scratch_url))
        try:
            async with async_sessionmaker(engine)() as session:
                await session.execute(
                    text(
                        "TRUNCATE check_run, criterion, rubric, manuscript, instructor "
                        "RESTART IDENTITY CASCADE"
                    )
                )
                await session.commit()
        finally:
            await engine.dispose()

    asyncio.run(_truncate())
    yield


@pytest.fixture(autouse=True)
def _logged_in(client, _clean_tables, api_scratch_url):
    """Seed the demo instructor and log the shared `client` in — every
    rubric route now requires auth (BUG-001/D-020). Depends on
    `_clean_tables` explicitly so the truncate always runs first."""
    import asyncio

    from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

    from app.db import sqlalchemy_url
    from app.models.instructor import Instructor

    async def seed():
        engine = create_async_engine(sqlalchemy_url(api_scratch_url))
        try:
            async with async_sessionmaker(engine, expire_on_commit=False)() as session:
                session.add(
                    Instructor(
                        email="prof@tip.edu.ph",
                        display_name="Prof",
                        password_hash=hash_password("s3cret!"),
                    )
                )
                await session.commit()
        finally:
            await engine.dispose()

    asyncio.run(seed())
    client.post("/auth/login", json={"email": "prof@tip.edu.ph", "password": "s3cret!"})


def _rubric_pdf(tmp_path):
    b = PdfBuilder()
    b.new_page().line("REQUIRED FORMAT CHECKLIST", bold=True)
    b.line("The manuscript must include an abstract.")
    return b.save(tmp_path / "rubric.pdf")


@pytest.mark.parametrize(
    ("filename", "payload", "unrecognized"),
    [
        ("picture.png", b"\x89PNG\r\n\x1a\nimage signature", True),
        ("picture.jpg", b"\xff\xd8\xff\xe0image signature", True),
        ("picture.pdf", b"\xff\xd8\xff\xe0image signature", True),
        ("unknown.docx", b"unrecognized content", True),
        ("damaged.pdf", b"%PDF-1.7\ntruncated", False),
        ("damaged.docx", b"PK\x03\x04truncated", False),
    ],
    ids=["png", "jpeg", "renamed-image", "unknown", "unreadable-pdf", "unreadable-zip"],
)
def test_upload_recovery_advice_and_failed_draft_cleanup(
    client, monkeypatch, filename, payload, unrecognized
):
    from app import messages
    from app.rubric import service

    decompose = AsyncMock(side_effect=AssertionError("Rejected bytes must not reach grading"))
    monkeypatch.setattr(service, "attempt_decomposition", decompose)
    response = client.post(
        "/rubrics", files={"file": (filename, payload, "application/octet-stream")}
    )
    assert response.status_code == 422, response.text
    expected = (
        "VERIDICAL could not recognize this file as a supported document. "
        "Choose a document in one of these formats: .docx, .pdf."
        if unrecognized
        else messages.FILE_UNREADABLE
    )
    assert response.json() == {"error": {"code": "file_malformed", "message": expected}}
    decompose.assert_not_awaited()
    families = client.get("/rubric-families")
    assert families.status_code == 200 and families.json() == []
    assert not list((get_settings().data_dir / "rubric_uploads").iterdir())


def test_post_rubrics_returns_persisted_criteria(client, tmp_path):
    path = _rubric_pdf(tmp_path)
    with path.open("rb") as fh:
        resp = client.post(
            "/rubrics",
            params={"title": "Demo Rubric"},
            files={"file": ("rubric.pdf", fh, "application/pdf")},
        )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["title"] == "Demo Rubric"
    assert body["version"] == 1
    assert body["is_active"] is False
    assert len(body["criteria"]) >= 1
    assert {c["type"] for c in body["criteria"]} <= {"structural", "semantic"}
    assert sum(c["weight"] for c in body["criteria"]) == pytest.approx(100.0)


def test_post_rubrics_accepts_a_pdf_whose_text_carries_nul(client, tmp_path):
    """BUG-550: glyph code 0 in the format's text failed the AI audit write
    (PostgreSQL JSONB cannot hold U+0000) and the upload returned 500."""
    from tests.test_document_nul_text import glyph_zero_pdf

    path = glyph_zero_pdf(tmp_path / "rubric.pdf", r"The manuscript must include an\000 abstract.")
    with path.open("rb") as fh:
        resp = client.post(
            "/rubrics",
            params={"title": "NUL Rubric"},
            files={"file": ("rubric.pdf", fh, "application/pdf")},
        )
    assert resp.status_code == 200, resp.text
    assert "\x00" not in resp.text and "\\u0000" not in resp.text


def _upload(client, tmp_path, *, title="Demo Rubric"):
    path = _rubric_pdf(tmp_path)
    with path.open("rb") as fh:
        resp = client.post(
            "/rubrics",
            params={"title": title},
            files={"file": ("rubric.pdf", fh, "application/pdf")},
        )
    assert resp.status_code == 200, resp.text
    return resp.json()


def test_get_rubric_returns_the_uploaded_rubric(client, tmp_path):
    uploaded = _upload(client, tmp_path)
    resp = client.get(f"/rubrics/{uploaded['id']}")
    assert resp.status_code == 200
    assert resp.json()["id"] == uploaded["id"]


def test_get_rubric_404s_on_an_unknown_id(client):
    resp = client.get("/rubrics/999999")
    assert resp.status_code == 404
    assert resp.json()["error"]["code"] == "not_found"


def test_put_criteria_edit_and_confirm_round_trips_through_http(client, tmp_path):
    uploaded = _upload(client, tmp_path)
    edited = [
        {**c, "weight": 60.0, "type": "semantic" if c["type"] == "structural" else "structural"}
        for c in uploaded["criteria"]
    ]
    body = _confirm(client, uploaded["id"], edited)
    assert body["is_active"] is True

    reloaded = client.get(f"/rubrics/{uploaded['id']}").json()
    assert reloaded["is_active"] is True
    assert all(c["weight"] == 60.0 for c in reloaded["criteria"])


@pytest.mark.parametrize("confirm", [False, True])
def test_put_replaced_criteria_returns_persisted_membership_and_reusable_ids(
    client, tmp_path, confirm
):
    """F2.3: the successful response must describe the rows just saved."""
    uploaded = _upload(client, tmp_path)
    old_ids = {row["id"] for row in uploaded["criteria"]}
    survivor = uploaded["criteria"][-1]
    levels = [
        {"level": 1, "name": "Limited", "descriptor": "Limited comparison", "points": 1},
        {"level": 2, "name": "Developed", "descriptor": "Developed comparison", "points": 3},
    ]
    replacement = [
        {
            "type": "semantic",
            "text": "Findings explain their relationship to earlier work",
            "evidence": "Discussion comparison",
            "weight": 7,
            "levels": levels,
        },
        {**survivor, "text": "Retained requirement after edit", "weight": 2},
    ]
    route = f"/rubrics/{uploaded['id']}/criteria"
    saved = client.put(route, json={"criteria": replacement})
    assert saved.status_code == 200, saved.text
    if confirm:
        # Hand-written levels are not the document's words: they are kept
        # only once the instructor says so (BUG-566).
        saved = client.put(
            route,
            json={
                "criteria": saved.json()["criteria"],
                "confirm": True,
                "acknowledgements": _acks(saved.json()),
            },
        )
        assert saved.status_code == 200, saved.text
    body = saved.json()
    assert body["is_active"] is confirm
    assert len(body["criteria"]) == len(replacement)
    assert [row["text"] for row in body["criteria"]] == [row["text"] for row in replacement]
    assert body["criteria"][0]["id"] not in old_ids
    assert body["criteria"][1]["id"] == survivor["id"]
    assert [row["position"] for row in body["criteria"]] == [0, 1]
    assert body["criteria"][0]["evidence"] == replacement[0]["evidence"]
    assert body["criteria"][0]["weight"] == 7
    assert [
        {k: v for k, v in level.items() if k != "source"} for level in body["criteria"][0]["levels"]
    ] == levels
    assert {level["source"]["match"] for level in body["criteria"][0]["levels"]} == {"unmatched"}
    fresh = client.get(f"/rubrics/{uploaded['id']}")
    assert fresh.status_code == 200
    assert body == fresh.json()
    again = client.put(route, json={"criteria": body["criteria"], "confirm": confirm})
    assert again.status_code == 200, again.text
    assert again.json()["revision"] == body["revision"] + 1
    assert {**again.json(), "revision": None} == {**body, "revision": None}


def test_put_criteria_rejects_an_empty_list(client, tmp_path):
    """Server-side backstop for 'deleting all criteria -> confirm disabled'
    (the frontend disables the button; the API also refuses to persist
    zero criteria as defense in depth)."""
    uploaded = _upload(client, tmp_path)
    resp = client.put(f"/rubrics/{uploaded['id']}/criteria", json={"criteria": [], "confirm": True})
    assert resp.status_code == 422


def test_put_criteria_rejects_a_zero_or_negative_weight_at_the_api(client, tmp_path):
    """BUG-052 fix item 1: 'validate only against degenerate input (zero,
    negative, all-zero); enforce at the API, not just the button.' A
    weight of 999 is a valid relative value under the DECIDED model
    (weights don't need to total 100) and must NOT be rejected -- only
    zero/negative, which can never mean anything as a relative weight."""
    uploaded = _upload(client, tmp_path)
    criteria = [{**c, "weight": 999.0} for c in uploaded["criteria"]]
    resp = client.put(
        f"/rubrics/{uploaded['id']}/criteria", json={"criteria": criteria, "confirm": False}
    )
    assert resp.status_code == 200, resp.text  # a large relative weight is legitimate

    zero = [{**c, "weight": 0} for c in uploaded["criteria"]]
    resp = client.put(
        f"/rubrics/{uploaded['id']}/criteria", json={"criteria": zero, "confirm": True}
    )
    assert resp.status_code == 422

    negative = [{**c, "weight": -5} for c in uploaded["criteria"]]
    resp = client.put(
        f"/rubrics/{uploaded['id']}/criteria", json={"criteria": negative, "confirm": True}
    )
    assert resp.status_code == 422

    # And the rubric must NOT have been activated by either rejected attempt.
    reloaded = client.get(f"/rubrics/{uploaded['id']}").json()
    assert reloaded["is_active"] is False


def _acks(rubric):
    """Keep every open parse issue as written (the fake model's drafts of
    these fixtures always leave one)."""
    return [
        {"issue_id": d["id"], "reason_code": "other", "note": "Checked in the test."}
        for d in rubric.get("parse_issue_details") or []
        if d["status"] == "open"
    ]


def _confirm(client, rubric_id, criteria):
    rubric = client.get(f"/rubrics/{rubric_id}").json()
    resp = client.put(
        f"/rubrics/{rubric_id}/criteria",
        json={"criteria": criteria, "confirm": True, "acknowledgements": _acks(rubric)},
    )
    assert resp.status_code == 200, resp.text
    return resp.json()


def test_confirm_waits_for_open_issues_but_keeps_the_save(client, tmp_path):
    """BUG-566: confirming with an open parse issue is refused; the edit made
    in the same request is still saved."""
    uploaded = _upload(client, tmp_path)
    assert uploaded["parse_status"] == "needs_review"
    open_ids = [d["id"] for d in uploaded["parse_issue_details"] if d["status"] == "open"]
    edited = [{**c, "weight": 3.0} for c in uploaded["criteria"]]
    route = f"/rubrics/{uploaded['id']}/criteria"
    resp = client.put(route, json={"criteria": edited, "confirm": True})
    assert resp.status_code == 409, resp.text
    details = resp.json()["error"]["details"]
    assert details["open_issue_ids"] == open_ids
    after = client.get(f"/rubrics/{uploaded['id']}").json()
    assert after["is_active"] is False
    assert {c["weight"] for c in after["criteria"]} == {after["criteria"][0]["weight"]}
    assert after["revision"] == uploaded["revision"] + 1
    # The refusal carries the saved revision, so the next save can be checked
    # against it without reloading.
    assert details["revision"] == after["revision"]
    again = client.put(
        route,
        json={
            "criteria": after["criteria"],
            "confirm": True,
            "acknowledgements": _acks(after),
            "expected_revision": details["revision"],
        },
    )
    assert again.status_code == 200, again.text
    assert again.json()["is_active"] is True


def test_keeping_a_flagged_wording_records_who_when_and_why(client, tmp_path):
    uploaded = _upload(client, tmp_path)
    issue = uploaded["parse_issue_details"][0]
    route = f"/rubrics/{uploaded['id']}/criteria"
    ack = {"issue_id": issue["id"], "reason_code": "compared_whole_format", "note": " As printed. "}
    resp = client.put(route, json={"criteria": uploaded["criteria"], "acknowledgements": [ack]})
    assert resp.status_code == 200, resp.text
    kept = next(d for d in resp.json()["parse_issue_details"] if d["id"] == issue["id"])
    assert kept["status"] == "acknowledged"
    resolution = kept["resolution"]
    assert resolution["kind"] == "acknowledged"
    assert resolution["reason_code"] == "compared_whole_format"
    assert resolution["note"] == "As printed."
    assert resolution["by_name"] and resolution["by_instructor_id"] and resolution["at"]
    unknown = {"issue_id": 999, "reason_code": "other", "note": None}
    missing = client.put(
        route, json={"criteria": uploaded["criteria"], "acknowledgements": [unknown]}
    )
    assert missing.status_code == 404


def test_a_hand_saved_single_level_does_not_break_the_recheck(client, tmp_path):
    uploaded = _upload(client, tmp_path)
    one_level = {
        "type": "semantic",
        "text": "Overall impression",
        "weight": 1,
        "levels": [{"level": 1, "name": "Met", "descriptor": "Meets it.", "points": 1}],
    }
    resp = client.put(
        f"/rubrics/{uploaded['id']}/criteria",
        json={"criteria": [*uploaded["criteria"], one_level]},
    )
    assert resp.status_code == 200, resp.text


def _whole_format_issue(rubric):
    issue = rubric["parse_issue_details"][0]
    assert issue["code"] == "other" and issue["criterion_id"] is None
    return issue


def test_a_whole_format_item_needs_a_note_saying_what_was_compared(client, tmp_path):
    """A whole-format item has no wording to keep: the note is the only record
    of what the instructor compared, so it is refused without one and nothing
    is saved."""
    uploaded = _upload(client, tmp_path)
    issue = _whole_format_issue(uploaded)
    route = f"/rubrics/{uploaded['id']}/criteria"
    bare = {"issue_id": issue["id"], "reason_code": "compared_whole_format", "note": None}
    refused = client.put(route, json={"criteria": uploaded["criteria"], "acknowledgements": [bare]})
    assert refused.status_code == 422, refused.text
    assert refused.json()["error"]["code"] == "acknowledgement_note_required"
    assert refused.json()["error"]["details"]["issue_ids"] == [issue["id"]]
    unchanged = client.get(f"/rubrics/{uploaded['id']}").json()
    assert unchanged["revision"] == uploaded["revision"]
    assert _whole_format_issue(unchanged)["status"] == "open"
    noted = {**bare, "note": "Read pages 1 to 3 against every criterion."}
    resp = client.put(route, json={"criteria": uploaded["criteria"], "acknowledgements": [noted]})
    assert resp.status_code == 200, resp.text
    kept = _whole_format_issue(resp.json())
    assert kept["status"] == "acknowledged"
    assert kept["resolution"]["reason_code"] == "compared_whole_format"
    assert kept["resolution"]["note"] == "Read pages 1 to 3 against every criterion."


def test_a_keep_reason_must_fit_the_item(client, tmp_path):
    """The whole-format reason is refused on a level, and a level's wording
    reasons are refused on a whole-format item, so the record never states
    the wrong reason."""
    uploaded = _upload(client, tmp_path)
    issue = _whole_format_issue(uploaded)
    route = f"/rubrics/{uploaded['id']}/criteria"
    wording = {"issue_id": issue["id"], "reason_code": "document_wording", "note": "As printed."}
    resp = client.put(route, json={"criteria": uploaded["criteria"], "acknowledgements": [wording]})
    assert resp.status_code == 422, resp.text
    assert resp.json()["error"]["code"] == "acknowledgement_reason_mismatch"
    assert resp.json()["error"]["details"]["issue_ids"] == [issue["id"]]
    explained = {**wording, "reason_code": "other"}
    resp = client.put(
        route, json={"criteria": uploaded["criteria"], "acknowledgements": [explained]}
    )
    assert resp.status_code == 200, resp.text


def test_a_kept_whole_format_item_opens_again_when_a_criterion_is_removed(client, tmp_path):
    """The item was checked against the criteria that existed then. Removing
    one reopens it; adding one cannot cover less of the format and does not."""
    uploaded = _upload(client, tmp_path)
    assert len(uploaded["criteria"]) >= 2
    issue = _whole_format_issue(uploaded)
    route = f"/rubrics/{uploaded['id']}/criteria"
    ack = {
        "issue_id": issue["id"],
        "reason_code": "compared_whole_format",
        "note": "Compared all pages.",
    }
    kept = client.put(route, json={"criteria": uploaded["criteria"], "acknowledgements": [ack]})
    assert kept.status_code == 200, kept.text
    assert _whole_format_issue(kept.json())["status"] == "acknowledged"

    added = {"type": "semantic", "text": "Delivery", "weight": 1}
    grown = client.put(route, json={"criteria": [*kept.json()["criteria"], added]})
    assert grown.status_code == 200, grown.text
    assert _whole_format_issue(grown.json())["status"] == "acknowledged"

    fewer = grown.json()["criteria"][1:]
    shrunk = client.put(route, json={"criteria": fewer})
    assert shrunk.status_code == 200, shrunk.text
    reopened = _whole_format_issue(shrunk.json())
    assert reopened["status"] == "open" and reopened["resolution"] is None
    assert reopened["id"] == issue["id"]


def test_an_active_format_keeps_its_criteria_until_a_reopened_item_is_checked_again(
    client, tmp_path
):
    """On an active format, removing a criterion that a kept whole-format item
    was checked against is refused and nothing is saved, unless the item is
    checked again in the same save."""
    uploaded = _upload(client, tmp_path)
    issue = _whole_format_issue(uploaded)
    active = _confirm(client, uploaded["id"], uploaded["criteria"])
    assert active["is_active"] is True
    route = f"/rubrics/{uploaded['id']}/criteria"
    fewer = active["criteria"][1:]
    refused = client.put(route, json={"criteria": fewer})
    assert refused.status_code == 409, refused.text
    assert issue["id"] in refused.json()["error"]["details"]["open_issue_ids"]
    unchanged = client.get(f"/rubrics/{uploaded['id']}").json()
    assert len(unchanged["criteria"]) == len(active["criteria"])
    assert _whole_format_issue(unchanged)["status"] == "acknowledged"
    again = {
        "issue_id": issue["id"],
        "reason_code": "compared_whole_format",
        "note": "Compared again without it.",
    }
    resp = client.put(route, json={"criteria": fewer, "acknowledgements": [again]})
    assert resp.status_code == 200, resp.text
    assert len(resp.json()["criteria"]) == len(fewer)
    assert _whole_format_issue(resp.json())["resolution"]["note"] == "Compared again without it."


def test_a_missing_keep_note_is_refused_in_the_standard_envelope():
    import asyncio
    import json

    from app.errors import AcknowledgementNoteRequiredError
    from app.main import veridical_error_handler

    resp = asyncio.run(
        veridical_error_handler(None, AcknowledgementNoteRequiredError("Add a note.", [4, 7]))
    )
    assert resp.status_code == 422
    assert json.loads(resp.body) == {
        "error": {
            "code": "acknowledgement_note_required",
            "message": "Add a note.",
            "details": {"issue_ids": [4, 7]},
        }
    }


def test_a_save_on_an_older_revision_is_refused(client, tmp_path):
    uploaded = _upload(client, tmp_path)
    route = f"/rubrics/{uploaded['id']}/criteria"
    first = client.put(
        route, json={"criteria": uploaded["criteria"], "expected_revision": uploaded["revision"]}
    )
    assert first.status_code == 200, first.text
    stale = client.put(
        route, json={"criteria": uploaded["criteria"], "expected_revision": uploaded["revision"]}
    )
    assert stale.status_code == 409
    assert stale.json()["error"]["code"] == "rubric_changed"


def test_reupload_with_family_id_creates_v2_over_http(client, tmp_path):
    v1 = _upload(client, tmp_path, title="V1")
    _confirm(client, v1["id"], v1["criteria"])

    path = _rubric_pdf(tmp_path)
    with path.open("rb") as fh:
        resp = client.post(
            "/rubrics",
            params={"title": "V2", "family_id": v1["rubric_family_id"]},
            files={"file": ("rubric.pdf", fh, "application/pdf")},
        )
    assert resp.status_code == 200, resp.text
    v2 = resp.json()
    assert v2["rubric_family_id"] == v1["rubric_family_id"]
    assert v2["version"] == 2
    assert v2["is_active"] is False

    v1_reloaded = client.get(f"/rubrics/{v1['id']}").json()
    assert v1_reloaded["version"] == 1
    assert v1_reloaded["is_latest_version"] is False  # superseded by v2


@pytest.mark.parametrize("payload", [b"unrecognized image content", b"%PDF-1.7\ntruncated"])
def test_rejected_version_preserves_active_family_criteria_and_source(
    client, tmp_path, monkeypatch, payload
):
    import hashlib

    import docx

    from app.rubric import service

    path = tmp_path / "required-format.docx"
    document = docx.Document()
    document.add_heading("REQUIRED FORMAT CHECKLIST", level=1)
    document.add_paragraph("The manuscript must include an abstract.")
    document.save(path)
    uploaded = client.post(
        "/rubrics", files={"file": (path.name, path.read_bytes(), "application/octet-stream")}
    )
    assert uploaded.status_code == 200, uploaded.text
    initial = uploaded.json()
    _confirm(client, initial["id"], initial["criteria"])
    before = client.get(f"/rubrics/{initial['id']}").json()
    assert before["is_active"] and before["is_latest_version"]
    versions_url = f"/rubric-families/{initial['rubric_family_id']}/versions"
    versions_before = client.get(versions_url).json()
    upload_dir = get_settings().data_dir / "rubric_uploads"
    source_before = {
        p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in upload_dir.iterdir()
    }
    assert len(source_before) == 1
    decompose = AsyncMock(side_effect=AssertionError("Rejected bytes must not reach grading"))
    monkeypatch.setattr(service, "attempt_decomposition", decompose)
    rejected = client.post(
        "/rubrics",
        params={"family_id": initial["rubric_family_id"]},
        files={"file": ("replacement.pdf", payload, "application/octet-stream")},
    )
    assert rejected.status_code == 422, rejected.text
    assert rejected.json()["error"]["code"] == "file_malformed"
    decompose.assert_not_awaited()
    assert client.get(f"/rubrics/{initial['id']}").json() == before
    assert client.get(versions_url).json() == versions_before
    assert {
        p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in upload_dir.iterdir()
    } == source_before


def test_activate_switches_which_version_is_active(client, tmp_path):
    v1 = _upload(client, tmp_path, title="V1")
    _confirm(client, v1["id"], v1["criteria"])
    path = _rubric_pdf(tmp_path)
    with path.open("rb") as fh:
        v2 = client.post(
            "/rubrics",
            params={"title": "V2", "family_id": v1["rubric_family_id"]},
            files={"file": ("rubric.pdf", fh, "application/pdf")},
        ).json()

    # BUG-566: a version with open parse issues waits for the instructor.
    blocked = client.post(f"/rubrics/{v2['id']}/activate")
    assert blocked.status_code == 409, blocked.text
    error = blocked.json()["error"]
    assert error["code"] == "unresolved_rubric_issues"
    assert error["details"]["open_issue_ids"] == [d["id"] for d in v2["parse_issue_details"]]
    kept = client.put(
        f"/rubrics/{v2['id']}/criteria",
        json={
            "criteria": v2["criteria"],
            "acknowledgements": _acks(v2),
            "expected_revision": v2["revision"],
        },
    )
    assert kept.status_code == 200, kept.text
    assert {d["status"] for d in kept.json()["parse_issue_details"]} == {"acknowledged"}

    resp = client.post(f"/rubrics/{v2['id']}/activate")
    assert resp.status_code == 200, resp.text
    assert resp.json()["is_active"] is True

    v1_after = client.get(f"/rubrics/{v1['id']}").json()
    assert v1_after["is_active"] is False


def test_delete_rubric_succeeds_when_no_reports_exist(client, tmp_path):
    uploaded = _upload(client, tmp_path)
    resp = client.delete(f"/rubrics/{uploaded['id']}")
    assert resp.status_code == 204
    assert client.get(f"/rubrics/{uploaded['id']}").status_code == 404


def test_list_rubric_families_and_versions(client, tmp_path):
    v1 = _upload(client, tmp_path, title="V1")
    _confirm(client, v1["id"], v1["criteria"])
    path = _rubric_pdf(tmp_path)
    with path.open("rb") as fh:
        client.post(
            "/rubrics",
            params={"title": "V2", "family_id": v1["rubric_family_id"]},
            files={"file": ("rubric.pdf", fh, "application/pdf")},
        )

    families = client.get("/rubric-families").json()
    assert len(families) == 1
    assert families[0]["rubric_family_id"] == v1["rubric_family_id"]
    assert families[0]["is_active"] is True  # v1, still active — v2 not confirmed yet

    versions = client.get(f"/rubric-families/{v1['rubric_family_id']}/versions").json()
    assert [v["version"] for v in versions] == [2, 1]


def test_citation_recency_setting_round_trips_and_carries_to_a_new_version(client, tmp_path):
    v1 = _upload(client, tmp_path, title="V1")
    assert v1["citation_recency"]["setting_sources"]["window_years"] == "default"
    family_url = f"/rubric-families/{v1['rubric_family_id']}/citation-recency"

    resp = client.put(family_url, json={"enabled": True, "window_years": 3})
    assert resp.status_code == 200, resp.text
    assert resp.json()["window_years"] == 3
    assert resp.json()["setting_sources"]["window_years"] == "format"
    assert client.get(f"/rubrics/{v1['id']}").json()["citation_recency"]["enabled"] is True

    for bad in ({"window_years": 0}, {"window_years": "3"}, {"basis": "x"}):
        assert client.put(family_url, json=bad).status_code == 422

    _confirm(client, v1["id"], v1["criteria"])
    path = _rubric_pdf(tmp_path)
    with path.open("rb") as fh:
        resp = client.post(
            "/rubrics",
            params={"title": "V2", "family_id": v1["rubric_family_id"]},
            files={"file": ("rubric.pdf", fh, "application/pdf")},
        )
    assert resp.status_code == 200, resp.text
    assert resp.json()["citation_recency"]["window_years"] == 3

    cleared = client.put(family_url, json={})
    assert cleared.json()["setting_sources"]["window_years"] == "default"


def test_citation_recency_setting_404s_for_an_unknown_family(client):
    resp = client.put(
        "/rubric-families/00000000-0000-0000-0000-000000000000/citation-recency",
        json={"enabled": True},
    )
    assert resp.status_code == 404


def test_a_kept_whole_format_item_opens_again_when_a_criterion_is_reworded(client, tmp_path):
    uploaded = _upload(client, tmp_path)
    issue = _whole_format_issue(uploaded)
    route = f"/rubrics/{uploaded['id']}/criteria"
    ack = {
        "issue_id": issue["id"],
        "reason_code": "compared_whole_format",
        "note": "Compared all pages.",
    }
    kept = client.put(route, json={"criteria": uploaded["criteria"], "acknowledgements": [ack]})
    assert kept.status_code == 200, kept.text
    # Saving again unchanged leaves it kept.
    same = client.put(route, json={"criteria": kept.json()["criteria"]})
    assert _whole_format_issue(same.json())["status"] == "acknowledged"
    reworded = [{**c} for c in same.json()["criteria"]]
    reworded[0]["text"] = "Something else entirely"
    resp = client.put(route, json={"criteria": reworded})
    assert resp.status_code == 200, resp.text
    assert _whole_format_issue(resp.json())["status"] == "open"


def _criterion(criterion_id, text):
    from app.models.rubric import Criterion

    return Criterion(id=criterion_id, rubric_id=1, text=text, levels=None)


def _kept_whole_format(resolution):
    return {
        "id": 1,
        "code": "other",
        "criterion_id": None,
        "level": None,
        "message": "the criteria cover too little of the format",
        "status": "acknowledged",
        "resolution": resolution,
    }


def test_an_item_kept_before_criteria_were_recorded_takes_the_current_criteria():
    """It is not reopened on every save (that would make an active format
    impossible to save); it is checked against the criteria from then on."""
    from app.models.rubric import Rubric
    from app.rubric.issues import reopen_whole_format

    resolution = {"kind": "acknowledged", "reason_code": "other", "note": "Old keep.", "at": "x"}
    rubric = Rubric(parse_issue_details=[_kept_whole_format(resolution)])
    rows = [_criterion(1, "Purpose"), _criterion(2, "Content")]
    reopen_whole_format(rubric, rows)
    detail = rubric.parse_issue_details[0]
    assert detail["status"] == "acknowledged"
    assert set(detail["resolution"]["criteria"]) == {"1", "2"}
    reopen_whole_format(rubric, rows[:1])
    assert rubric.parse_issue_details[0]["status"] == "open"


def test_the_whole_format_reason_is_refused_on_a_level():
    from app.errors import AcknowledgementReasonMismatchError
    from app.models.instructor import Instructor
    from app.models.rubric import Rubric
    from app.rubric.issues import acknowledge
    from app.rubric.schemas import IssueAcknowledgement

    level_issue = {
        "id": 3,
        "code": "joins_cells",
        "criterion_id": 1,
        "level": 2,
        "message": "m",
        "status": "open",
        "resolution": None,
    }
    rubric = Rubric(parse_issue_details=[level_issue])
    ack = IssueAcknowledgement(
        issue_id=3, reason_code="compared_whole_format", note="Compared the whole original."
    )

    class _Session:
        def add(self, _row):
            raise AssertionError("nothing is recorded for a refused keep")

    with pytest.raises(AcknowledgementReasonMismatchError) as caught:
        acknowledge(
            _Session(),
            rubric,
            [ack],
            Instructor(id=1, display_name="Ana"),
            [_criterion(1, "Purpose")],
            get_settings(),
        )
    assert caught.value.details == {"issue_ids": [3]}
    assert rubric.parse_issue_details[0]["status"] == "open"
