"""V-081/F8: replay exact saved source images, never reinterpret a document on read."""

import hashlib
import io
import zipfile
from datetime import UTC, datetime
from types import SimpleNamespace

import docx
import pytest

from app.config import get_settings
from app.errors import ConflictError, GoneError, NotFoundError
from app.ingest.docx import extract_document
from app.ingest.service import _write_raw_store, raw_store_path
from app.report.docx_images import _verify_occurrence, manuscript_docx_image_for
from app.report.service import manuscript_paragraphs_for
from tests.test_docx_source_media import _jpeg_bytes, _source
from tests.test_ingest_docx import _png_bytes


@pytest.fixture()
def saved(tmp_path):
    settings = get_settings().model_copy(update={"data_dir": tmp_path, "storage_backend": "local"})
    path = tmp_path / "source.docx"
    _source(path)
    extraction = extract_document(str(path), settings)
    manuscript = SimpleNamespace(
        id=1,
        file_ref=str(path),
        purged_at=None,
        content_hash=extraction.docx_source_sha256,
        section_tree=extraction.section_tree.model_dump(),
    )
    _write_raw_store(extraction, raw_store_path(settings, manuscript.id))
    return settings, manuscript, extraction, path


def _save(saved):
    settings, manuscript, extraction, _ = saved
    _write_raw_store(extraction, raw_store_path(settings, manuscript.id))


def test_native_table_and_images_survive_saved_json_and_public_projection(saved):
    settings, manuscript, _, _ = saved
    response = manuscript_paragraphs_for(manuscript, settings)
    assert [row.paragraph for row in response.paragraphs] == [0, 2]
    assert response.tables[0].paragraph == 3
    assert response.tables[0].rows == [["Original table text", ""]]
    assert [(row.image_index, row.paragraph, row.available) for row in response.images] == [
        (0, 1, True),
        (1, 3, True),
    ]
    assert "relationship_id" not in response.model_dump_json()
    assert "word/media" not in response.model_dump_json()
    assert "sha256" not in response.model_dump_json()
    for index in (0, 1):
        assert manuscript_docx_image_for(manuscript, index, settings) == (_png_bytes(), "image/png")


def test_source_bound_images_and_table_ignore_misleading_pdf_name(saved):
    settings, manuscript, _, path = saved
    renamed = path.with_suffix(".pdf")
    path.rename(renamed)
    manuscript.file_ref = str(renamed)
    assert manuscript_docx_image_for(manuscript, 0, settings) == (_png_bytes(), "image/png")
    projection = manuscript_paragraphs_for(manuscript, settings)
    assert projection.tables[0].rows == [["Original table text", ""]]
    assert len(projection.images) == 2


@pytest.mark.parametrize("index", [-1, 2, 999999])
def test_out_of_range_index_is_not_python_negative_indexing(saved, index):
    settings, manuscript, _, _ = saved
    with pytest.raises(NotFoundError):
        manuscript_docx_image_for(manuscript, index, settings)


def test_purged_source_is_unavailable_even_if_original_bytes_remain(saved):
    settings, manuscript, _, _ = saved
    manuscript.purged_at = datetime.now(UTC)
    with pytest.raises(GoneError, match="purged"):
        manuscript_docx_image_for(manuscript, 0, settings)


def test_non_docx_source_is_not_opened_as_a_package(saved):
    from tests.test_ingest_fixtures import FIXTURE_DIR

    settings, manuscript, _, _ = saved
    manuscript.file_ref = str(FIXTURE_DIR / "native.pdf")
    with pytest.raises(ConflictError):
        manuscript_docx_image_for(manuscript, 0, settings)


@pytest.mark.parametrize("missing", ["source_hash", "descriptor"])
def test_legacy_record_does_not_reparse_or_fetch_original(saved, monkeypatch, missing):
    import app.report.docx_images as media

    settings, manuscript, extraction, _ = saved
    if missing == "source_hash":
        extraction.docx_source_sha256 = None
    else:
        extraction.images[0].docx_source = None
    _save(saved)
    monkeypatch.setattr(media, "manuscript_source_file", lambda *args: pytest.fail("Read original"))
    with pytest.raises(GoneError):
        manuscript_docx_image_for(manuscript, 0, settings)
    response = manuscript_paragraphs_for(manuscript, settings)
    assert response.images[0].available is False
    assert response.images[0].unavailable_reason == "source_not_recorded"


def test_replaced_cached_original_is_not_served_under_old_provenance(saved):
    settings, manuscript, _, path = saved
    path.write_bytes(path.read_bytes() + b"different source")
    with pytest.raises(GoneError):
        manuscript_docx_image_for(manuscript, 0, settings)


@pytest.mark.parametrize("corruption", ["descriptor", "paragraph"])
def test_valid_image_cannot_be_relocated_to_another_source_occurrence(saved, corruption):
    import io

    settings, manuscript, _, path = saved
    document = docx.Document(path)
    document.add_picture(io.BytesIO(_jpeg_bytes()))
    document.save(path)
    extraction = extract_document(str(path), settings)
    manuscript.content_hash = extraction.docx_source_sha256
    if corruption == "descriptor":
        extraction.images[0].docx_source = extraction.images[2].docx_source
    else:
        extraction.images[0].paragraph = extraction.images[2].paragraph
    _save((settings, manuscript, extraction, path))
    with pytest.raises(GoneError):
        manuscript_docx_image_for(manuscript, 0, settings)


@pytest.mark.parametrize(
    "field,value",
    [
        ("sha256", "0" * 64),
        ("byte_length", 1),
        ("width_px", 9),
        ("height_px", 9),
        ("content_type", "image/jpeg"),
        ("part_name", "word/media/absent.png"),
        ("relationship_id", "missing"),
    ],
)
def test_corrupt_saved_descriptor_is_not_used_to_serve_other_bytes(saved, field, value):
    settings, manuscript, extraction, _ = saved
    setattr(extraction.images[0].docx_source, field, value)
    _save(saved)
    with pytest.raises(GoneError):
        manuscript_docx_image_for(manuscript, 0, settings)


@pytest.mark.parametrize(
    "field,value",
    [("docx_image_max_bytes", 1), ("docx_image_max_pixels", 1), ("docx_image_max_count", 1)],
)
def test_current_limits_are_rechecked_when_serving_saved_media(saved, field, value):
    settings, manuscript, _, _ = saved
    settings = settings.model_copy(update={field: value})
    exception = NotFoundError if field == "docx_image_max_count" else GoneError
    with pytest.raises(exception):
        manuscript_docx_image_for(manuscript, 1, settings)


def _replace_relationships(saved, transform):
    """Alter only a synthetic fixture, then bind the fixture ledger to its new source."""
    settings, manuscript, extraction, path = saved
    with zipfile.ZipFile(path) as package:
        members = {entry.filename: package.read(entry) for entry in package.infolist()}
    name = "word/_rels/document.xml.rels"
    members[name] = transform(members[name])
    with zipfile.ZipFile(path, "w") as package:
        for member, content in members.items():
            package.writestr(member, content)
    manuscript.content_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    extraction.docx_source_sha256 = manuscript.content_hash
    _save(saved)


@pytest.mark.parametrize(
    "transform",
    [
        lambda xml: xml.replace(
            b'Target="media/image1.png"', b'Target="https://invalid.example/a"'
        ),
        lambda xml: xml.replace(
            b'Target="media/image1.png"', b'TargetMode="External" Target="media/image1.png"'
        ),
        lambda xml: xml.replace(b"relationships/image", b"relationships/hyperlink"),
        lambda xml: b"<broken",
    ],
)
def test_invalid_or_external_relationship_is_never_followed(saved, transform):
    settings, manuscript, _, _ = saved
    _replace_relationships(saved, transform)
    with pytest.raises(GoneError):
        manuscript_docx_image_for(manuscript, 0, settings)


def test_duplicate_package_member_is_not_selected_arbitrarily(saved):
    settings, manuscript, extraction, path = saved
    with (
        zipfile.ZipFile(path, "a") as package,
        pytest.warns(UserWarning, match="Duplicate name"),
    ):
        package.writestr("word/media/image1.png", _png_bytes())
    manuscript.content_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    extraction.docx_source_sha256 = manuscript.content_hash
    _save(saved)
    with pytest.raises(GoneError):
        manuscript_docx_image_for(manuscript, 0, settings)


def _occurrence_xml(image, *, leading=b"", trailing=b"", declaration=b""):
    relationship = image.docx_source.relationship_id.encode()
    return (
        declaration
        + b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
        b' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
        b' xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        b"<w:body>"
        + leading
        + b'<w:p/><w:p><w:drawing><a:blip r:embed="'
        + relationship
        + b'"/></w:drawing></w:p>'
        + trailing
        + b"</w:body></w:document>"
    )


def test_occurrence_lookup_stops_before_reading_large_trailing_body(saved):
    settings, _, extraction, _ = saved
    image = extraction.images[0]
    content = _occurrence_xml(image, trailing=b"<w:p/>" * 20_000)
    stream = io.BytesIO(content)
    _verify_occurrence(stream, 0, image, settings)
    assert stream.tell() < len(content) // 2


@pytest.mark.parametrize(
    "field,value",
    [("docx_image_source_xml_max_bytes", 100), ("docx_image_source_xml_max_nodes", 5)],
)
def test_occurrence_lookup_has_independent_byte_and_node_limits(saved, field, value):
    settings, _, extraction, _ = saved
    image = extraction.images[0]
    settings = settings.model_copy(update={field: value})
    stream = io.BytesIO(_occurrence_xml(image))
    with pytest.raises(GoneError):
        _verify_occurrence(stream, 0, image, settings)
    if field.endswith("bytes"):
        assert stream.tell() <= value + 1


def test_large_prefix_is_not_scanned_unboundedly_to_find_later_image(saved):
    settings, _, extraction, _ = saved
    image = extraction.images[0]
    content = _occurrence_xml(image, leading=b"<w:p/>" * 20_000)
    image = image.model_copy(update={"paragraph": image.paragraph + 20_000})
    settings = settings.model_copy(update={"docx_image_source_xml_max_nodes": 100})
    stream = io.BytesIO(content)
    with pytest.raises(GoneError):
        _verify_occurrence(stream, 0, image, settings)
    assert stream.tell() < len(content) // 2


def test_streaming_occurrence_keeps_actual_byte_limit_inclusive(saved):
    settings, _, extraction, _ = saved
    image = extraction.images[0]
    content = _occurrence_xml(image)
    settings = settings.model_copy(update={"docx_image_source_xml_max_bytes": len(content)})
    _verify_occurrence(io.BytesIO(content), 0, image, settings)


def test_occurrence_lookup_does_not_treat_revision_content_as_a_body_item(saved):
    settings, _, extraction, _ = saved
    image = extraction.images[0]
    content = _occurrence_xml(image, leading=b"<w:ins><w:p><w:drawing/></w:p></w:ins>")
    _verify_occurrence(io.BytesIO(content), 0, image, settings)


def test_occurrence_lookup_cannot_continue_into_a_second_document_body(saved):
    settings, _, extraction, _ = saved
    image = extraction.images[0]
    content = _occurrence_xml(image).replace(b"<w:body>", b"<w:body><w:p/></w:body><w:body>")
    with pytest.raises(GoneError):
        _verify_occurrence(io.BytesIO(content), 0, image, settings)
