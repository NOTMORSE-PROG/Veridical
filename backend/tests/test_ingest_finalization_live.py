"""F1: real transaction failures must not strand ordinary preparation."""

import asyncio
import hashlib
import os
import shutil
import socket
from pathlib import Path

import pytest
from docx import Document
from sqlalchemy import select, text
from sqlalchemy.exc import DBAPIError
from sqlalchemy.sql.dml import Delete, Update

from app.config import Settings
from app.errors import FileTooLargeError
from app.ingest import service
from app.ingest.schemas import ExtractionResult
from app.models.citation import Citation
from app.models.enums import CitationParseStatus, IngestFailureReason, IngestStatus
from app.models.instructor import Instructor
from app.models.manuscript import Manuscript

pytestmark = pytest.mark.skipif("DATABASE_URL" not in os.environ, reason="needs PostgreSQL")
SCRATCH_DB = "veridical_ingestfinaltest"
SAFE_LATE_FAILURE = "VERIDICAL could not finish preparing this manuscript. Try uploading it again."


@pytest.fixture(scope="module")
def scratch_url():
    from alembic import command
    from tests.test_schema import _admin_execute, _alembic_config, _swap_db

    base = os.environ["DATABASE_URL"]
    asyncio.run(_admin_execute(base, f'DROP DATABASE IF EXISTS "{SCRATCH_DB}"'))
    asyncio.run(_admin_execute(base, f'CREATE DATABASE "{SCRATCH_DB}"'))
    url = _swap_db(base, SCRATCH_DB)
    command.upgrade(_alembic_config(url), "head")
    yield url
    asyncio.run(_admin_execute(base, f'DROP DATABASE IF EXISTS "{SCRATCH_DB}" WITH (FORCE)'))


@pytest.fixture
async def session_factory(scratch_url):
    from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

    from app.db import sqlalchemy_url

    engine = create_async_engine(sqlalchemy_url(scratch_url))
    yield async_sessionmaker(engine, expire_on_commit=False)
    await engine.dispose()


@pytest.fixture(autouse=True)
async def clean_and_keyless(session_factory, monkeypatch):
    async with session_factory() as session:
        await session.execute(text("TRUNCATE manuscript, instructor RESTART IDENTITY CASCADE"))
        await session.commit()

    async def no_model(*_args, **_kwargs):
        raise AssertionError("These source controls must not request a model")

    monkeypatch.setattr(service, "get_llm_client_for", no_model)
    connect = socket.socket.connect
    connect_ex = socket.socket.connect_ex

    def local_only(original):
        def guarded(sock, address):
            if isinstance(address, tuple) and address[0] in {"127.0.0.1", "::1"}:
                return original(sock, address)
            raise AssertionError("Ingest finalization controls must not use external services")

        return guarded

    monkeypatch.setattr(socket.socket, "connect", local_only(connect))
    monkeypatch.setattr(socket.socket, "connect_ex", local_only(connect_ex))


async def seed(session_factory, tmp_path, suffix):
    fixture = Path(__file__).parent / "fixtures" / "ingest" / f"native.{suffix}"
    original_hash = hashlib.sha256(fixture.read_bytes()).hexdigest()
    source = tmp_path / f"source.{suffix}"
    shutil.copyfile(fixture, source)
    if suffix == "docx":
        document = Document(source)
        document.add_heading("References", 1)
        document.add_paragraph("Reyes, J. (2023). A Study. Journal of Examples.")
        document.save(source)
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    async with session_factory() as session:
        owner = Instructor(email="final-owner@example.test", display_name="Owner")
        other = Instructor(email="final-other@example.test", display_name="Other")
        session.add_all([owner, other])
        await session.flush()
        manuscript = Manuscript(
            instructor_id=owner.id,
            group_label="Late preparation control",
            original_filename=source.name,
            file_ref=str(source),
            content_hash=source_hash,
            ingest_status=IngestStatus.done,
            ingest_failure_reason=IngestFailureReason.file_too_large,
            ingest_failure_detail="Prior safe size-limit message",
            section_tree={"source": "none", "nodes": []},
        )
        session.add(manuscript)
        await session.flush()
        session.add(
            Citation(
                manuscript_id=manuscript.id,
                order_index=0,
                raw_text="Original, A. (2020). Old Reference. Example Journal.",
                authors=["Original, A."],
                year=2020,
                title="Old Reference",
                parse_status=CitationParseStatus.parsed,
            )
        )
        await session.commit()
        return manuscript.id, owner.id, other.id, source, source_hash, fixture, original_hash


async def snapshot(session_factory, manuscript_id):
    async with session_factory() as session:
        row = await session.get(Manuscript, manuscript_id)
        citations = list(
            await session.scalars(
                select(Citation)
                .where(Citation.manuscript_id == manuscript_id)
                .order_by(Citation.id)
            )
        )
        return {
            "id": row.id,
            "owner": row.instructor_id,
            "file_ref": row.file_ref,
            "content_hash": row.content_hash,
            "status": row.ingest_status,
            "reason": row.ingest_failure_reason,
            "detail": row.ingest_failure_detail,
            "section_tree": row.section_tree,
            "citations": [(item.id, item.order_index, item.raw_text) for item in citations],
        }


@pytest.mark.parametrize("existing", [None, "mismatch"])
async def test_fresh_docx_ingestion_binds_only_a_missing_source_identity(
    session_factory, tmp_path, existing
):
    from app.ingest.analytical_source import analytical_source, bind_analytical_source

    identity, _, _, source, expected_hash, _, _ = await seed(session_factory, tmp_path, "docx")
    previous = None if existing is None else "0" * 64
    async with session_factory() as session:
        manuscript = await session.get(Manuscript, identity)
        manuscript.content_hash = previous
        await session.commit()
        extraction = await service.ingest_manuscript(
            session,
            manuscript,
            source,
            Settings(_env_file=None, data_dir=tmp_path / "data"),
            read_images=False,
        )
        assert extraction.docx_source_sha256 == expected_hash
        assert manuscript.content_hash == (expected_hash if previous is None else previous)
        bound = analytical_source(bind_analytical_source(extraction, manuscript.content_hash))
        assert bound.state == ("complete" if previous is None else "unavailable")
    saved = await snapshot(session_factory, identity)
    assert saved["content_hash"] == (expected_hash if previous is None else previous)
    assert saved["status"] == IngestStatus.done


async def test_fresh_docx_identity_assignment_rolls_back_with_failed_finalization(
    session_factory, tmp_path
):
    identity, _, _, source, _, _, _ = await seed(session_factory, tmp_path, "docx")
    async with session_factory() as session:
        manuscript = await session.get(Manuscript, identity)
        manuscript.content_hash = None
        await session.commit()
    await install_fault(session_factory, identity, "deferred_commit")
    try:
        async with session_factory() as session:
            manuscript = await session.get(Manuscript, identity)
            with pytest.raises(DBAPIError):
                await service.ingest_manuscript(
                    session,
                    manuscript,
                    source,
                    Settings(_env_file=None, data_dir=tmp_path / "data"),
                    read_images=False,
                )
    finally:
        async with session_factory() as session:
            await session.execute(text("DROP TRIGGER ingest_finalize_test_fault ON manuscript"))
            await session.commit()
    saved = await snapshot(session_factory, identity)
    assert saved["content_hash"] is None
    assert saved["status"] == IngestStatus.failed
    assert saved["citations"][0][2] == "Original, A. (2020). Old Reference. Example Journal."


async def test_fresh_docx_identity_is_not_published_by_an_intermediate_consumer_commit(
    session_factory, tmp_path, monkeypatch
):
    identity, _, _, source, _, _, _ = await seed(session_factory, tmp_path, "docx")
    async with session_factory() as session:
        manuscript = await session.get(Manuscript, identity)
        manuscript.content_hash = None
        await session.commit()

    async def consumer_commits_then_fails(active_session, *_args):
        await active_session.commit()
        raise RuntimeError("Controlled optional-consumer failure")

    monkeypatch.setattr(
        service.vision, "prepare_image_readings", lambda *_args, **_kwargs: [object()]
    )
    monkeypatch.setattr(service, "get_llm_client_for", consumer_commits_then_fails)
    async with session_factory() as session:
        manuscript = await session.get(Manuscript, identity)
        with pytest.raises(RuntimeError, match="Controlled optional-consumer failure"):
            await service.ingest_manuscript(
                session, manuscript, source, Settings(_env_file=None, data_dir=tmp_path / "data")
            )
    saved = await snapshot(session_factory, identity)
    assert saved["content_hash"] is None
    assert saved["status"] == IngestStatus.failed


async def test_a_preparation_handed_off_before_finalizing_publishes_nothing(
    session_factory, tmp_path
):
    """An upload retired by another worker while this one still prepared it:
    the refusal raised inside the final transaction must leave the record
    neither ready nor rewritten as a preparation failure."""
    identity, _, _, source, _, _, _ = await seed(session_factory, tmp_path, "docx")
    calls = []

    async def handed_off(active_session):
        calls.append(active_session)
        raise service.PreparationHandedOff

    async with session_factory() as session:
        manuscript = await session.get(Manuscript, identity)
        with pytest.raises(service.PreparationHandedOff):
            await service.ingest_manuscript(
                session,
                manuscript,
                source,
                Settings(_env_file=None, data_dir=tmp_path / "data"),
                read_images=False,
                before_finalize=handed_off,
            )
        assert calls == [session]
    saved = await snapshot(session_factory, identity)
    assert saved["status"] == IngestStatus.processing
    assert saved["detail"] == "Prior safe size-limit message"
    assert saved["citations"][0][2] == "Original, A. (2020). Old Reference. Example Journal."


async def install_fault(session_factory, manuscript_id, fault):
    event = "UPDATE" if fault == "deferred_commit" else "INSERT"
    table = "manuscript" if fault == "deferred_commit" else "citation"
    predicate = (
        f"NEW.id = {manuscript_id} AND NEW.ingest_status = 'done'"
        if fault == "deferred_commit"
        else f"NEW.manuscript_id = {manuscript_id}"
    )
    timing = "AFTER" if fault == "deferred_commit" else "BEFORE"
    deferred = "DEFERRABLE INITIALLY DEFERRED" if fault == "deferred_commit" else ""
    constraint = "CONSTRAINT" if fault == "deferred_commit" else ""
    async with session_factory() as session:
        await session.execute(
            text("""
                CREATE OR REPLACE FUNCTION ingest_finalize_test_fault() RETURNS trigger
                LANGUAGE plpgsql AS $$ BEGIN
                  RAISE EXCEPTION 'private controlled database rejection' USING ERRCODE = '23514';
                  RETURN NEW;
                END $$
            """)
        )
        await session.execute(
            text(f"""CREATE {constraint} TRIGGER ingest_finalize_test_fault
                     {timing} {event} ON {table} {deferred}
                     FOR EACH ROW WHEN ({predicate})
                     EXECUTE FUNCTION ingest_finalize_test_fault()""")
        )
        await session.commit()
    return table


@pytest.mark.parametrize("suffix", ["pdf", "docx"])
@pytest.mark.parametrize("fault", ["after_delete", "citation_insert", "deferred_commit"])
async def test_late_database_failure_is_terminal_atomic_and_private(
    session_factory, tmp_path, monkeypatch, suffix, fault
):
    mid, owner, other, source, source_hash, fixture, fixture_hash = await seed(
        session_factory, tmp_path, suffix
    )
    before = await snapshot(session_factory, mid)
    settings = Settings(_env_file=None, data_dir=tmp_path, storage_backend="local")
    trigger_table = None
    if fault != "after_delete":
        trigger_table = await install_fault(session_factory, mid, fault)
    try:
        async with session_factory() as session:
            manuscript = await session.get(Manuscript, mid)
            execute = session.execute
            fired = False

            async def execute_with_fault(statement, *args, **kwargs):
                nonlocal fired
                result = await execute(statement, *args, **kwargs)
                if fault == "after_delete" and not fired and isinstance(statement, Delete):
                    fired = True
                    await execute(text("SELECT 1 / 0"))
                return result

            monkeypatch.setattr(session, "execute", execute_with_fault)
            with pytest.raises(DBAPIError) as caught:
                await service.ingest_manuscript(
                    session, manuscript, source, settings, read_images=False
                )
            assert caught.value.orig.sqlstate == ("22012" if fault == "after_delete" else "23514")
            # A separate reader observes recovery before the caller rolls back.
            after = await snapshot(session_factory, mid)
            assert after["status"] == IngestStatus.failed
            assert after["reason"] == IngestFailureReason.extraction_failed
            assert after["detail"] == SAFE_LATE_FAILURE
            assert await session.scalar(text("SELECT 1")) == 1
            for key in ("id", "owner", "file_ref", "content_hash", "section_tree", "citations"):
                assert after[key] == before[key]
        async with session_factory() as observer:
            own = await service.list_manuscripts(observer, owner)
            foreign = await service.list_manuscripts(observer, other)
            assert own.total == 1 and own.items[0].ingest_status == "failed"
            assert own.items[0].ingest_failure_detail == SAFE_LATE_FAILURE
            assert foreign.total == 0
        raw = service.raw_store_path(settings, mid)
        extraction = ExtractionResult.model_validate_json(raw.read_text(encoding="utf-8"))
        assert extraction.text_chars > 0
        assert hashlib.sha256(source.read_bytes()).hexdigest() == source_hash
        assert hashlib.sha256(fixture.read_bytes()).hexdigest() == fixture_hash
    finally:
        if trigger_table:
            async with session_factory() as cleanup:
                await cleanup.execute(
                    text(f"DROP TRIGGER ingest_finalize_test_fault ON {trigger_table}")
                )
                await cleanup.commit()


@pytest.mark.parametrize("suffix", ["pdf", "docx"])
async def test_successful_preparation_commits_new_citations_and_section_tree(
    session_factory, tmp_path, suffix
):
    mid, _owner, _other, source, source_hash, _fixture, _fixture_hash = await seed(
        session_factory, tmp_path, suffix
    )
    before = await snapshot(session_factory, mid)
    settings = Settings(_env_file=None, data_dir=tmp_path, storage_backend="local")
    async with session_factory() as session:
        manuscript = await session.get(Manuscript, mid)
        result = await service.ingest_manuscript(
            session, manuscript, source, settings, read_images=False
        )
    after = await snapshot(session_factory, mid)
    assert after["status"] == IngestStatus.done
    assert after["section_tree"] == result.section_tree.model_dump()
    assert after["citations"] and after["citations"] != before["citations"]
    assert after["content_hash"] == source_hash


async def test_early_page_limit_updates_the_original_cached_manuscript(session_factory, tmp_path):
    mid, _owner, _other, source, _source_hash, _fixture, _fixture_hash = await seed(
        session_factory, tmp_path, "pdf"
    )
    async with session_factory() as session:
        manuscript = await session.get(Manuscript, mid)
        with pytest.raises(FileTooLargeError):
            await service.ingest_manuscript(
                session,
                manuscript,
                source,
                Settings(
                    _env_file=None, data_dir=tmp_path, storage_backend="local", ingest_max_pages=3
                ),
                read_images=False,
            )
        # No intervening SELECT or caller refresh may be needed for correctness.
        assert manuscript.ingest_status == IngestStatus.failed
        assert manuscript.ingest_failure_reason == IngestFailureReason.file_too_large
        assert "6 pages" in manuscript.ingest_failure_detail
        assert "3 page limit" in manuscript.ingest_failure_detail
        after = await snapshot(session_factory, mid)
        assert after["status"] == manuscript.ingest_status
        assert after["detail"] == manuscript.ingest_failure_detail


@pytest.mark.parametrize("recovery_fault", ["rollback", "failure_write"])
async def test_recovery_failure_keeps_original_database_error_and_safe_logs(
    session_factory, tmp_path, monkeypatch, caplog, recovery_fault
):
    mid, _owner, _other, source, _source_hash, _fixture, _fixture_hash = await seed(
        session_factory, tmp_path, "docx"
    )
    original_errors = []
    private_marker = "private-database-diagnostic-must-not-be-logged"
    async with session_factory() as session:
        manuscript = await session.get(Manuscript, mid)
        execute = session.execute

        async def execute_with_fault(statement, *args, **kwargs):
            if isinstance(statement, Update) and recovery_fault == "failure_write":
                raise RuntimeError(private_marker)
            result = await execute(statement, *args, **kwargs)
            if isinstance(statement, Delete):
                try:
                    await execute(text("SELECT 1 / 0"))
                except DBAPIError as error:
                    original_errors.append(error)
                    raise
            return result

        async def unavailable_rollback():
            raise RuntimeError(private_marker)

        monkeypatch.setattr(session, "execute", execute_with_fault)
        if recovery_fault == "rollback":
            monkeypatch.setattr(session, "rollback", unavailable_rollback)
        with pytest.raises(DBAPIError) as caught:
            await service.ingest_manuscript(
                session,
                manuscript,
                source,
                Settings(_env_file=None, data_dir=tmp_path, storage_backend="local"),
                read_images=False,
            )
        assert caught.value is original_errors[0]
        assert caught.value.orig.sqlstate == "22012"
        assert private_marker not in caplog.text
        assert "Could not record manuscript preparation failure" in caplog.text
    after = await snapshot(session_factory, mid)
    # Recovery itself was unavailable; do not invent a durable terminal state.
    assert after["status"] == IngestStatus.processing
    assert after["detail"] == "Prior safe size-limit message"


@pytest.mark.parametrize("removal", ["delete", "purge"])
async def test_failure_recovery_does_not_restore_removed_manuscript_metadata(
    session_factory, tmp_path, monkeypatch, removal
):
    from sqlalchemy import delete, func, update

    mid, _owner, _other, source, _source_hash, _fixture, _fixture_hash = await seed(
        session_factory, tmp_path, "docx"
    )
    async with session_factory() as session:
        manuscript = await session.get(Manuscript, mid)
        execute, rollback = session.execute, session.rollback
        removed = False

        async def execute_with_fault(statement, *args, **kwargs):
            result = await execute(statement, *args, **kwargs)
            if isinstance(statement, Delete):
                await execute(text("SELECT 1 / 0"))
            return result

        async def rollback_then_remove():
            nonlocal removed
            await rollback()
            if not removed:
                removed = True
                async with session_factory() as other_session:
                    statement = (
                        delete(Manuscript).where(Manuscript.id == mid)
                        if removal == "delete"
                        else update(Manuscript)
                        .where(Manuscript.id == mid)
                        .values(purged_at=func.clock_timestamp())
                    )
                    await other_session.execute(statement)
                    await other_session.commit()

        monkeypatch.setattr(session, "execute", execute_with_fault)
        monkeypatch.setattr(session, "rollback", rollback_then_remove)
        with pytest.raises(DBAPIError):
            await service.ingest_manuscript(
                session,
                manuscript,
                source,
                Settings(_env_file=None, data_dir=tmp_path, storage_backend="local"),
                read_images=False,
            )
    async with session_factory() as observer:
        row = await observer.get(Manuscript, mid)
        if removal == "delete":
            assert row is None
        else:
            assert row.purged_at is not None
            assert row.ingest_status == IngestStatus.processing
            assert row.ingest_failure_detail == "Prior safe size-limit message"
