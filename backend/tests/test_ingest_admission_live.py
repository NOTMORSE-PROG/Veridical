"""F1/F2: cross-instructor upload admission precedes actual durable bookkeeping."""

import asyncio
import threading
from contextlib import suppress
from io import BytesIO

import pytest
from docx import Document
from sqlalchemy import func, select

from app.config import Settings
from app.errors import DocumentPreparationBusyError, FileTooLargeError
from app.ingest import admission, service
from app.library.schemas import LibrarySubmissionIn
from app.library.service import submit_library_entry
from app.models.group import Group
from app.models.instructor import Instructor
from app.models.library import LibraryEntry
from app.models.manuscript import Manuscript
from app.models.rubric import Rubric
from app.rubric.service import create_rubric_from_upload
from tests.test_ingest_finalization_live import clean_and_keyless as clean_and_keyless
from tests.test_ingest_finalization_live import pytestmark as pytestmark
from tests.test_ingest_finalization_live import scratch_url as scratch_url
from tests.test_ingest_finalization_live import session_factory as session_factory


class NoModel:
    async def complete(self, *args, **kwargs):
        pytest.fail("Busy rejection must not invoke a model")


async def owner_counts(session, owner_id):
    counts = {}
    for model in (Manuscript, Group, Rubric):
        counts[model.__tablename__] = await session.scalar(
            select(func.count()).select_from(model).where(model.instructor_id == owner_id)
        )
    counts["library_entry"] = await session.scalar(
        select(func.count())
        .select_from(LibraryEntry)
        .join(Manuscript)
        .where(Manuscript.instructor_id == owner_id)
    )
    return counts


@pytest.mark.parametrize("timeout", [False, True])
async def test_other_instructor_uploads_cannot_bypass_surviving_native_work(
    session_factory, tmp_path, monkeypatch, timeout
):
    document = Document()
    document.add_paragraph("Synthetic preparation admission control.")
    document.add_heading("References", 1)
    document.add_paragraph("Reyes, J. (2024). A synthetic source. Journal of Examples.")
    buffer = BytesIO()
    document.save(buffer)
    data = buffer.getvalue()
    original_path = tmp_path / "first.docx"
    original_path.write_bytes(data)
    settings = Settings(
        _env_file=None,
        storage_backend="local",
        data_dir=tmp_path / "data",
        ingest_extraction_timeout_seconds=0.025 if timeout else 10,
    )
    started, release = threading.Event(), threading.Event()
    lock = threading.Lock()
    counts = {"active": 0, "peak": 0, "entered": 0}
    original_select = service.select_extractor

    def controlled_select(suffix):
        native = original_select(suffix)

        def extract(*args, **kwargs):
            with lock:
                counts["entered"] += 1
                position = counts["entered"]
                counts["active"] += 1
                counts["peak"] = max(counts["peak"], counts["active"])
            try:
                if position == 1:
                    started.set()
                    assert release.wait(10), "Every test native barrier must be released"
                return native(*args, **kwargs)
            finally:
                with lock:
                    counts["active"] -= 1

        return extract

    monkeypatch.setattr(service, "select_extractor", controlled_select)
    first_task = None
    runtime = admission.get_preparation_admission(settings)
    try:
        async with session_factory() as first_session, session_factory() as second_session:
            first_owner = Instructor(email="first@example.test", display_name="Synthetic first")
            second_owner = Instructor(email="second@example.test", display_name="Synthetic second")
            first_session.add_all([first_owner, second_owner])
            await first_session.commit()
            first = Manuscript(
                instructor_id=first_owner.id, group_label="First", file_ref=str(original_path)
            )
            first_session.add(first)
            await first_session.commit()
            first_task = asyncio.create_task(
                service.ingest_manuscript(
                    first_session, first, original_path, settings, read_images=False
                )
            )
            assert await asyncio.to_thread(started.wait, 5)
            if timeout:
                with pytest.raises(FileTooLargeError):
                    await first_task
                await first_session.refresh(first)
                assert first.ingest_status.value == "failed"
            assert counts["active"] == 1 and runtime.active_owners == 1
            before = await owner_counts(second_session, second_owner.id)
            assert all(value == 0 for value in before.values())
            files_before = sorted(
                path.relative_to(settings.data_dir)
                for path in settings.data_dir.rglob("*")
                if path.is_file()
            )
            consumed = []

            async def chunks():
                consumed.append(True)
                yield data

            for attempt in (
                lambda: service.ingest_upload(
                    second_session,
                    chunks(),
                    "second.docx",
                    "Second",
                    settings,
                    instructor_id=second_owner.id,
                ),
                lambda: create_rubric_from_upload(
                    second_session,
                    chunks(),
                    "format.docx",
                    "Format",
                    NoModel(),
                    settings,
                    instructor_id=second_owner.id,
                ),
                lambda: submit_library_entry(
                    second_session,
                    chunks(),
                    "source.docx",
                    LibrarySubmissionIn(),
                    instructor_id=second_owner.id,
                    settings=settings,
                ),
            ):
                with pytest.raises(DocumentPreparationBusyError) as rejected:
                    await attempt()
                assert rejected.value.code == "document_preparation_busy"
                assert await owner_counts(second_session, second_owner.id) == before
                assert consumed == [] and counts["entered"] == 1
                assert (
                    sorted(
                        path.relative_to(settings.data_dir)
                        for path in settings.data_dir.rglob("*")
                        if path.is_file()
                    )
                    == files_before
                )
            release.set()
            if not timeout:
                await first_task
                await first_session.refresh(first)
                assert first.ingest_status.value == "done"

            async def wait_for_drain():
                while not runtime.drained:
                    await asyncio.sleep(0.005)

            await asyncio.wait_for(wait_for_drain(), 5)
            retry_settings = settings.model_copy(update={"ingest_extraction_timeout_seconds": 10})
            accepted, result, reference_count, _, _ = await service.ingest_upload(
                second_session,
                chunks(),
                "second.docx",
                "Explicit retry",
                retry_settings,
                instructor_id=second_owner.id,
            )
            assert accepted.ingest_status.value == "done" and reference_count == 1
            assert result.anchor_kind == "paragraph"
            assert counts == {"active": 0, "peak": 1, "entered": 2}
            assert len(consumed) == 1 and runtime.drained
            assert original_path.read_bytes() == data
    finally:
        release.set()
        if first_task is not None and not first_task.done():
            with suppress(Exception, asyncio.CancelledError):
                await first_task

        async def cleanup_drain():
            while not runtime.drained:
                await asyncio.sleep(0.005)

        await asyncio.wait_for(cleanup_drain(), 10)
