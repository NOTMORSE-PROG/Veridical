"""F1/F8: source-bound native images are display data, not image assessment."""

import hashlib
import io
import zipfile
from pathlib import Path

import docx
import pymupdf
import pytest
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn

from app.config import get_settings
from app.errors import FileMalformedError
from app.ingest.docx import extract_document
from app.ingest.patterns import load_patterns
from app.ingest.schemas import ExtractionResult
from app.ingest.vision import select_images
from tests.test_ingest_docx import _png_bytes


def _source(path: Path):
    document = docx.Document()
    document.add_heading("Original heading", level=1)
    paragraph = document.add_paragraph()
    picture = paragraph.add_run().add_picture(io.BytesIO(_png_bytes()))
    picture._inline.docPr.set("descr", "Original author-supplied image description")
    document.add_paragraph("Figure 1. Original caption.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Original table text"
    table.cell(0, 1).paragraphs[0].add_run().add_picture(io.BytesIO(_png_bytes()))
    document.save(path)
    return document


def test_native_images_keep_original_bytes_identity_and_body_locations(tmp_path):
    path = tmp_path / "source.docx"
    _source(path)
    result = extract_document(str(path), get_settings())
    assert result.docx_source_sha256 == hashlib.sha256(path.read_bytes()).hexdigest()
    assert [image.paragraph for image in result.images] == [1, 3]
    assert all(image.page is None and image.bbox is None for image in result.images)
    assert result.images[0].source_alt_text == "Original author-supplied image description"
    for image in result.images:
        source = image.docx_source
        assert source is not None
        assert source.sha256 == hashlib.sha256(_png_bytes()).hexdigest()
        assert source.content_type == "image/png"
        assert (source.width_px, source.height_px) == (8, 8)
        assert source.byte_length == len(_png_bytes())
        assert image.display_unavailable_reason is None
    assert [block.paragraph for block in result.blocks] == [0, 2]
    assert result.blocks[-1].text == "Figure 1. Original caption."
    assert result.tables[0].paragraph == 3
    assert result.tables[0].rows == [["Original table text", ""]]
    restored = ExtractionResult.model_validate_json(result.model_dump_json())
    assert restored == result
    assert (
        select_images(restored, load_patterns(get_settings().ingest_patterns_file), get_settings())
        == []
    )


@pytest.mark.parametrize("linked_attribute", ["link", "embed"])
def test_external_picture_relationship_is_not_fetched_or_displayed(tmp_path, linked_attribute):
    path = tmp_path / "external.docx"
    document = _source(path)
    paragraph = document.paragraphs[1]
    blip = paragraph._p.xpath(".//a:blip")[0]
    blip.attrib.clear()
    rid = document.part.relate_to(
        "https://invalid.example/private-image", RT.IMAGE, is_external=True
    )
    blip.set(qn(f"r:{linked_attribute}"), rid)
    document.save(path)
    result = extract_document(str(path), get_settings())
    assert result.images[0].docx_source is None
    assert result.images[0].display_unavailable_reason == "external_image"


@pytest.mark.parametrize(
    "setting,value",
    [("docx_image_max_bytes", 1), ("docx_image_max_pixels", 1)],
)
def test_display_limits_do_not_discard_other_native_content(tmp_path, setting, value):
    path = tmp_path / "bounded.docx"
    _source(path)
    settings = get_settings().model_copy(update={setting: value})
    result = extract_document(str(path), settings)
    assert result.images[0].docx_source is None
    assert result.images[0].display_unavailable_reason == "image_limit"
    assert result.tables[0].rows[0][0] == "Original table text"
    assert result.blocks[-1].text == "Figure 1. Original caption."


def test_duplicate_package_members_are_rejected_before_source_mapping(tmp_path):
    path = tmp_path / "ambiguous.docx"
    _source(path)
    with zipfile.ZipFile(path, "a") as package:
        member = package.read("word/document.xml")
        with pytest.warns(UserWarning, match="Duplicate name"):
            package.writestr("word/document.xml", member)
    with pytest.raises(FileMalformedError):
        extract_document(str(path), get_settings())


def test_legacy_image_inventory_has_no_invented_display_provenance():
    from app.ingest.schemas import ImageBlock

    legacy = ImageBlock.model_validate({"paragraph": 2})
    assert legacy.docx_source is None
    assert legacy.source_alt_text is None


def _jpeg_bytes():
    return pymupdf.Pixmap(_png_bytes()).tobytes("jpeg")


@pytest.mark.parametrize("kind", ["png", "jpeg"])
def test_excessive_metadata_segments_are_rejected_before_header_object_allocation(
    kind, monkeypatch
):
    import app.ingest.docx_media as media

    settings = get_settings().model_copy(update={"docx_image_max_metadata_segments": 16})
    if kind == "png":
        original = _png_bytes()
        media.raster_header(original, "image/png", settings)
        # Zero-length ancillary chunks; synthetic input, not a source manuscript.
        blob = original[:-12] + (b"\x00\x00\x00\x00teSt\x00\x00\x00\x00" * 17) + original[-12:]
    else:
        original = _jpeg_bytes()
        media.raster_header(original, "image/jpeg", settings)
        sos = original.index(b"\xff\xda")
        blob = original[:sos] + (b"\xff\xfe\x00\x02" * 17) + original[sos:]
    monkeypatch.setattr(media.Image, "from_blob", lambda blob: pytest.fail("Allocated metadata"))
    with pytest.raises(ValueError, match="image_limit"):
        media.raster_header(blob, f"image/{kind}", settings)


@pytest.mark.parametrize("length", [30, 50, 100])
def test_truncated_jpeg_is_an_unavailable_image_not_an_uncaught_parser_exception(length):
    from app.ingest.docx_media import raster_header

    with pytest.raises(ValueError, match="invalid_image"):
        raster_header(_jpeg_bytes()[:length], "image/jpeg", get_settings())


def test_rejected_repeated_parts_are_inspected_once(tmp_path, monkeypatch):
    import app.ingest.docx_media as media

    path = tmp_path / "repeated.docx"
    _source(path)
    calls = []

    def reject(*args):
        calls.append(1)
        raise ValueError("invalid_image")

    monkeypatch.setattr(media, "raster_header", reject)
    result = extract_document(str(path), get_settings())
    assert len(calls) == 1
    assert [image.display_unavailable_reason for image in result.images] == [
        "invalid_image",
        "invalid_image",
    ]
    assert result.tables[0].rows[0][0] == "Original table text"
