"""Native Library activation persists and projects the actual bounded preparation."""

import asyncio
import json
import os
from io import BytesIO

import pytest
from docx import Document
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.checks.reuse.query import _archive_scope
from app.config import get_settings
from app.db import sqlalchemy_url
from tests.test_library_router_live import (
    OWNER,
    STRANGER,
    _fetchval,
    _login,
    _submit,
)
from tests.test_library_router_live import (
    api_scratch_url as api_scratch_url,
)
from tests.test_library_router_live import (
    client as client,
)

pytestmark = pytest.mark.skipif("DATABASE_URL" not in os.environ, reason="needs live PostgreSQL")


def _word(comment):
    document = Document()
    document.add_heading("Chapter 1: Introduction", 1)
    for i in range(15):
        document.add_paragraph(
            f"Observation {i} supports a distinct software design with source evidence."
        )
    document.core_properties.comments = comment
    out = BytesIO()
    document.save(out)
    return out.getvalue()


def test_partial_add_duplicate_decision_reload_and_orphan_scope(
    client, api_scratch_url, monkeypatch
):
    monkeypatch.setenv("REUSE_MAX_PASSAGES_PER_MANUSCRIPT", "3")
    monkeypatch.setenv("REUSE_PASSAGE_CHUNK_WORDS", "5")
    get_settings.cache_clear()

    _login(client, OWNER)
    added = _submit(client, _word("first"), "indexing.docx", title="Private indexed title")
    assert added["outcome"] == "added", added
    item = added["item"]
    mid = item["manuscript_id"]
    scope = item["passage_indexing"]
    assert scope["state"] == "partial" and scope["passages_indexed"] == 3
    assert scope["passages_found"] > 3
    assert (
        _fetchval(
            api_scratch_url,
            "SELECT count(*) FROM manuscript_passage_archive WHERE manuscript_id=$1",
            mid,
        )
        == 3
    )
    audit = json.loads(
        _fetchval(
            api_scratch_url,
            "SELECT payload::text FROM audit_log "
            "WHERE event_type='library_entry_added' AND manuscript_id=$1",
            mid,
        )
    )
    assert audit["passages_found"] == scope["passages_found"]
    assert audit["passages_embedded"] == audit["passages_stored"] == 3
    assert audit["passage_indexing"]["passages_found"] == scope["passages_found"]
    assert client.get(f"/library/{mid}").json()["passage_indexing"] == scope
    with monkeypatch.context() as altered:
        altered.setenv("EMBEDDING_MODEL_ID", "not-used-for-read-only-projection")
        altered.setenv("REUSE_MAX_PASSAGES_PER_MANUSCRIPT", "1")
        get_settings.cache_clear()
        assert client.get(f"/library/{mid}").json()["passage_indexing"] == scope
    get_settings.cache_clear()
    waiting = _submit(client, _word("second"), "different.docx")
    assert waiting["outcome"] == "possible_duplicate", waiting
    other = waiting["item"]["manuscript_id"]
    decision = client.post(f"/library/{other}/decision", json={"action": "add"})
    assert decision.status_code == 200, decision.text
    assert decision.json()["passage_indexing"] == scope
    client.post("/auth/logout")
    _login(client, STRANGER)
    foreign = client.get(f"/library/{mid}").json()
    assert foreign["passage_indexing"] == scope and foreign["title"] is None
    assert "source_sha256" not in json.dumps(foreign)

    async def aggregate():
        engine = create_async_engine(sqlalchemy_url(api_scratch_url))
        try:
            async with async_sessionmaker(engine)() as session:
                from sqlalchemy import text

                await session.execute(
                    text("DELETE FROM library_entry WHERE manuscript_id=:id"), {"id": other}
                )
                await session.commit()
                total, saved = await _archive_scope(
                    session, exclude_manuscript_ids=set(), settings=get_settings()
                )
                assert total == 2 and saved["partial"] and saved["not_recorded"]
                total, saved = await _archive_scope(
                    session, exclude_manuscript_ids={mid}, settings=get_settings()
                )
                assert total == 1 and not saved["partial"] and saved["not_recorded"]
        finally:
            await engine.dispose()

    asyncio.run(aggregate())
    get_settings.cache_clear()


def test_sql_scope_matches_strict_projection_for_stored_jsonb(client, api_scratch_url):
    from sqlalchemy import select, update

    from app.library.passage_indexing import project_passage_indexing
    from app.models.library import LibraryEntry
    from app.models.manuscript import Manuscript, ManuscriptArchive

    _login(client, OWNER)
    item = _submit(client, _word("oracle"), "oracle.docx")["item"]
    mid = item["manuscript_id"]

    async def exercise():
        engine = create_async_engine(sqlalchemy_url(api_scratch_url))
        try:
            async with async_sessionmaker(engine)() as session:
                source_hash, model_id, snapshot = (
                    await session.execute(
                        select(
                            Manuscript.content_hash,
                            ManuscriptArchive.model_id,
                            LibraryEntry.passage_indexing_snapshot,
                        )
                        .join(ManuscriptArchive, ManuscriptArchive.manuscript_id == Manuscript.id)
                        .join(LibraryEntry, LibraryEntry.manuscript_id == Manuscript.id)
                        .where(Manuscript.id == mid)
                    )
                ).one()
                base = snapshot | {"limit": 3, "passages_found": 7, "passages_indexed": 3}
                cases = [
                    None,
                    {},
                    [],
                    True,
                    "v1",
                    base,
                    base | {"passages_found": 0, "passages_indexed": 0},
                    base | {"limit": 0, "passages_indexed": 0},
                    base
                    | {
                        "inventory_state": "unavailable",
                        "passages_found": None,
                        "passages_indexed": 0,
                    },
                    base | {"source_sha256": "b" * 64},
                    base | {"model_id": "different"},
                    base | {"passages_found": True},
                    base | {"passages_found": "7"},
                    base | {"passages_found": 7.0},
                    base | {"passages_indexed": 4},
                    base | {"limit": -1},
                    base | {"limit": 3.0},
                    base | {"limit": 10**30, "passages_indexed": 7},
                    base | {"version": "v0"},
                ]
                missing = base | {"inventory_state": "unavailable", "passages_indexed": 0}
                del missing["passages_found"]
                cases.append(missing)
                for value in cases:
                    await session.execute(
                        update(LibraryEntry)
                        .where(LibraryEntry.manuscript_id == mid)
                        .values(passage_indexing_snapshot=value)
                    )
                    await session.commit()
                    stored = await session.scalar(
                        select(LibraryEntry.passage_indexing_snapshot).where(
                            LibraryEntry.manuscript_id == mid
                        )
                    )
                    expected = project_passage_indexing(stored, source_hash, model_id).state
                    total, summary = await _archive_scope(
                        session, exclude_manuscript_ids=set(), settings=get_settings()
                    )
                    assert total == 1
                    assert summary == dict(
                        version="v1",
                        partial=expected == "partial",
                        unavailable=expected == "unavailable",
                        not_recorded=expected == "not_recorded",
                    ), value
        finally:
            await engine.dispose()

    asyncio.run(exercise())
