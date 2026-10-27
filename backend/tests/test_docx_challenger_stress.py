"""Adversarial challenger stress tests for DOCX pipeline & ReportLab export.

Written by challenger_m2_rem_2 to empirically stress-test:
1. Native DOCX table evaluation in tables.py (genuine ratio, zero/partial captions, anchor fallback)
2. Manuscript.file_format resolution with unusual paths, extensions, and heuristic SectionTrees
3. Multi-page ReportLab PDF export with high volume, varied outcomes, and layout invariants
"""

import tempfile
from pathlib import Path

import docx as docx_lib
import fitz

from app.checks.rules import RuleContext, get_rule
from app.checks.rules.tables import RULE_ID as TABLE_FORMAT_RULE_ID
from app.config import get_settings
from app.ingest.docx import extract_document
from app.ingest.schemas import SectionTree, TableBlock
from app.ingest.service import _write_raw_store, raw_store_path
from app.models.enums import ResultOutcome
from app.models.manuscript import Manuscript
from app.report.export import build_report_pdf
from app.report.schemas import (
    CriterionResultOut,
    FlagExportDetailOut,
    FlagSummaryOut,
    PageEvidenceImageOut,
    ReportExportData,
    ReportOut,
    ResolutionOut,
)
from app.report.service import manuscript_paragraphs_for


class MockCriterion:
    def __init__(self, text: str):
        self.text = text


def test_tables_ratio_authentic_zero_and_partial():
    """Verify that table_caption_min_ratio never returns a fake 1.0 or passed outcome
    when tables are uncaptioned or under-captioned, across various counts.
    """
    rule = get_rule(TABLE_FORMAT_RULE_ID)
    criterion = MockCriterion("All tables must be properly formatted with descriptive captions")

    # Scenario 1: Zero tables in DOCX manuscript with images present
    # Must fail honestly, NOT return not_applicable
    ctx_no_tables = RuleContext(
        manuscript_id=10,
        anchor_kind="paragraph",
        page_count=0,
        section_tree=SectionTree(source="styles", nodes=[]),
        geometry=[],
        tables=[],
        vision_status="none",
        has_images=True,
        citations=[],
        margin_tolerance_pts=5.0,
        citation_style_min_ratio=0.8,
        table_caption_min_ratio=0.8,
    )
    outcome = rule.run(criterion, ctx_no_tables)
    assert outcome.outcome == ResultOutcome.failed
    assert "No tables were found" in outcome.detail["reason"]
    assert outcome.anchor == "document"

    # Scenario 2: N tables with 0 captions (caption=None and caption="")
    for n in (1, 2, 5, 10):
        tables = [
            TableBlock(
                paragraph=i * 3,
                rows=[["c1", "c2"], ["v1", "v2"]],
                caption=None if i % 2 == 0 else "",
            )
            for i in range(n)
        ]
        ctx_uncaptioned = RuleContext(
            manuscript_id=11,
            anchor_kind="paragraph",
            page_count=0,
            section_tree=SectionTree(source="styles", nodes=[]),
            geometry=[],
            tables=tables,
            vision_status="none",
            has_images=False,
            citations=[],
            margin_tolerance_pts=5.0,
            citation_style_min_ratio=0.8,
            table_caption_min_ratio=0.8,
        )
        res = rule.run(criterion, ctx_uncaptioned)
        assert res.outcome == ResultOutcome.failed
        assert res.detail["total_tables"] == n
        assert res.detail["captioned"] == 0
        assert res.detail["ratio"] == 0.0
        assert res.anchor == "paragraph 0"

    # Scenario 3: Under-threshold ratio (e.g. 3 of 5 captioned = 0.6 < 0.8)
    tables_mixed = [
        TableBlock(paragraph=0, rows=[["a", "b"]], caption="Table 1: First"),
        TableBlock(paragraph=3, rows=[["c", "d"]], caption="Table 2: Second"),
        TableBlock(paragraph=6, rows=[["e", "f"]], caption="Table 3: Third"),
        TableBlock(paragraph=9, rows=[["g", "h"]], caption=None),
        TableBlock(paragraph=12, rows=[["i", "j"]], caption=None),
    ]
    ctx_under = RuleContext(
        manuscript_id=12,
        anchor_kind="paragraph",
        page_count=0,
        section_tree=SectionTree(source="styles", nodes=[]),
        geometry=[],
        tables=tables_mixed,
        vision_status="none",
        has_images=False,
        citations=[],
        margin_tolerance_pts=5.0,
        citation_style_min_ratio=0.8,
        table_caption_min_ratio=0.8,
    )
    res_under = rule.run(criterion, ctx_under)
    assert res_under.outcome == ResultOutcome.failed
    assert res_under.detail["total_tables"] == 5
    assert res_under.detail["captioned"] == 3
    assert abs(res_under.detail["ratio"] - 0.6) < 1e-6
    assert res_under.anchor == "paragraph 9"

    # Scenario 4: Meets threshold (4 of 5 captioned = 0.8 >= 0.8)
    tables_mixed[3].caption = "Table 4: Fourth"
    res_pass = rule.run(criterion, ctx_under)
    assert res_pass.outcome == ResultOutcome.passed
    assert res_pass.detail["total_tables"] == 5
    assert res_pass.detail["captioned"] == 4
    assert abs(res_pass.detail["ratio"] - 0.8) < 1e-6
    assert res_pass.anchor == "paragraph 0"

    # Scenario 5: Missing paragraph and page anchor fallback
    table_no_anchor = [TableBlock(paragraph=None, page=None, rows=[["a"]], caption=None)]
    ctx_no_anchor = RuleContext(
        manuscript_id=13,
        anchor_kind="paragraph",
        page_count=0,
        section_tree=SectionTree(source="styles", nodes=[]),
        geometry=[],
        tables=table_no_anchor,
        vision_status="none",
        has_images=False,
        citations=[],
        margin_tolerance_pts=5.0,
        citation_style_min_ratio=0.8,
        table_caption_min_ratio=0.8,
    )
    res_no_anchor = rule.run(criterion, ctx_no_anchor)
    assert res_no_anchor.anchor == "document"


def test_manuscript_file_format_adversarial_inputs():
    """Verify Manuscript.file_format resolution with URL queries, multiple dots,
    unusual paths, and nested heuristic SectionTrees.
    """
    # 1. URL queries and fragments
    m1 = Manuscript(
        id=1,
        original_filename="thesis_final.docx?token=xyz123&action=download",
        file_ref="storage/bucket/1.docx",
    )
    assert m1.file_format == "docx"

    # 2. Multiple dots and spaces
    m2 = Manuscript(
        id=2,
        original_filename="T.I.P. Capstone Draft.v2.final.docx",
        file_ref="storage/bucket/2",
    )
    assert m2.file_format == "docx"

    # 3. Uppercase extension
    m3 = Manuscript(
        id=3,
        original_filename="MANUSCRIPT.DOCX",
        file_ref=None,
    )
    assert m3.file_format == "docx"

    # 4. Storage blob without extension, but PK\x03\x04 on disk
    with tempfile.NamedTemporaryFile(suffix="", delete=False) as tf:
        tf.write(b"PK\x03\x04\x14\x00\x06\x00")
        tf_path = tf.name

    try:
        m4 = Manuscript(
            id=4,
            original_filename="upload_stream_raw",
            file_ref=tf_path,
        )
        assert m4.file_format == "docx"
    finally:
        Path(tf_path).unlink(missing_ok=True)

    # 5. Remote/missing path, but SectionTree has heuristic paragraph nodes
    tree_dict = {
        "source": "heuristics",
        "nodes": [
            {
                "title": "CHAPTER I",
                "paragraph": 0,
                "children": [
                    {
                        "title": "Background of the Study",
                        "paragraph": 4,
                        "children": [
                            {"title": "Objectives", "paragraph": 10, "children": []}
                        ],
                    }
                ],
            }
        ],
    }
    m5 = Manuscript(
        id=5,
        original_filename="student_doc",
        file_ref="s3://remote-bucket/nonexistent-key",
        section_tree=tree_dict,
    )
    assert m5.file_format == "docx"

    # 6. Deeply nested SectionTree where top nodes have no paragraph, but leaf has paragraph
    tree_deep = {
        "source": "heuristics",
        "nodes": [
            {
                "title": "Part A",
                "children": [
                    {
                        "title": "Section 1",
                        "children": [
                            {"title": "Subsection 1.1", "paragraph": 15, "children": []}
                        ],
                    }
                ],
            }
        ],
    }
    m6 = Manuscript(
        id=6,
        original_filename="file_without_ext",
        file_ref="s3://remote/file",
        section_tree=tree_deep,
    )
    assert m6.file_format == "docx"


def test_manuscript_paragraphs_end_to_end_heuristic_docx():
    """Build a real DOCX file with heuristic headings, unusual storage path,
    save raw store, and test manuscript_paragraphs_for().
    """
    settings = get_settings()

    doc = docx_lib.Document()
    # Heading with bold styling (heuristic, not Word styles)
    p0 = doc.add_paragraph()
    r0 = p0.add_run("CHAPTER 1: PROBLEM STATEMENT")
    r0.bold = True
    doc.add_paragraph("This is the introductory text of the problem statement.")
    p2 = doc.add_paragraph()
    r2 = p2.add_run("CHAPTER 2: REVIEW OF RELATED LITERATURE")
    r2.bold = True
    doc.add_paragraph("Literature review text with relevant citations.")

    temp_dir = Path(tempfile.mkdtemp())
    docx_file = temp_dir / "raw_storage_object_uuid_7777"
    doc.save(str(docx_file))

    ext = extract_document(docx_file, settings)
    assert ext.anchor_kind == "paragraph"
    assert len(ext.blocks) >= 2

    m_id = 777777
    store_p = raw_store_path(settings, m_id)
    store_p.parent.mkdir(parents=True, exist_ok=True)
    _write_raw_store(ext, store_p)

    m = Manuscript(
        id=m_id,
        original_filename="student_submission_no_ext",
        file_ref=str(docx_file),
        section_tree=ext.section_tree.model_dump(),
        content_hash="hash_777777",
    )

    assert m.file_format == "docx"

    paragraphs_out = manuscript_paragraphs_for(m, settings)
    assert len(paragraphs_out.paragraphs) >= 2
    # Check that text is properly returned
    full_text = " ".join(p.text for p in paragraphs_out.paragraphs)
    assert "CHAPTER 1: PROBLEM STATEMENT" in full_text
    assert "Literature review text" in full_text


def test_docx_reportlab_multipage_export_comprehensive_stress():
    """Stress test multi-page ReportLab PDF export:
    - 50 criteria across structural and semantic categories
    - Both passed and failed criteria
    - Pending review criteria
    - Criteria resolved by instructor judgment
    - Multiple flags with docx_layout page image markers
    - Decision recorded with multiline reasoning note
    - Verifies page count >= 4, running headers, footers, and zero text block clipping
    """
    results = []
    # 50 criteria
    for i in range(50):
        is_pending = (i % 10 == 7)
        is_resolved = (i % 10 == 8)
        outcome = "passed" if (i % 3 != 0 and not is_pending) else "failed"
        if is_pending:
            outcome = "needs_review"
        resolution = (
            ResolutionOut(
                type="override",
                reason="Overridden to pass with instructor validation.",
                ai_majority_verdict="failed",
            )
            if is_resolved
            else None
        )

        results.append(
            CriterionResultOut(
                check_result_id=i + 1,
                criterion_id=i + 1,
                text=f"Criterion {i + 1}: Comprehensive Defense Readiness Evaluation Requirement.",
                type="structural" if i % 2 == 0 else "semantic",
                weight=1.0,
                weight_importance="high" if i % 4 == 0 else "med",
                kind="rule" if i % 2 == 0 else "ai",
                outcome=outcome,
                score=100.0 if outcome == "passed" else 0.0,
                basis="Rule-based validation" if i % 2 == 0 else "Rubric criteria grading",
                anchor=f"paragraph {i * 3}",
                reasoning=(
                    f"Detailed evaluation notes for criterion {i + 1} with academic context. "
                    "Ensures manuscript follows rigorous capstone formatting guidelines."
                ),
                reason=(
                    "Criterion satisfied"
                    if outcome == "passed"
                    else "Formatting issues observed."
                ),
                evidence=[],
                resolution=resolution,
            )
        )

    flags = [
        FlagSummaryOut(
            id=201 + j,
            check_kind="citations" if j % 2 == 0 else "internal_agreement",
            severity="high" if j == 0 else "medium",
            criterion_text=f"Integrity Check Criterion {j + 1}",
            evidence_excerpt=f"Identified anomaly at paragraph {j * 6 + 2}: unmatched claim.",
            page_anchor=f"paragraph {j * 6 + 2}",
            overridden=(j == 1),
            affects_readiness=True,
        )
        for j in range(5)
    ]

    flag_details = {
        f.id: FlagExportDetailOut(
            flag_id=f.id,
            finding_headline=f"Integrity Finding {f.id}",
            check_kind=f.check_kind,
            category="citation_anomaly",
            severity=f.severity,
            page_anchor=f.page_anchor,
            affects_readiness=f.affects_readiness,
            excerpt=f.evidence_excerpt,
            location_label=f.page_anchor,
            reasoning=f"Automated integrity scanner identified possible issue at {f.page_anchor}.",
            override_reason=(
                "Instructor verified reference exists in appendix." if f.overridden else None
            ),
        )
        for f in flags
    }

    page_images = {
        f.id: PageEvidenceImageOut(
            image_png=None,
            page=None,
            reason="DOCX layout - no page image",
            reason_code="docx_layout",
        )
        for f in flags
    }

    report = ReportOut(
        check_run_id=999,
        manuscript_id=888,
        manuscript_group_label="BSIT Capstone Group 12 - Defense Readiness",
        manuscript_original_filename="group12_capstone_final_submission.docx",
        rubric_title="TIP Manila BSIT Capstone Defense Rubric 2026",
        status="conditionally_ready",
        composite_score=82.0,
        thresholds={"ready_min_score": 85.0, "not_ready_max_score": 60.0},
        reason="Conditionally ready: 5 criteria require instructor review before final defense.",
        flag_deduction=5.0,
        unresolved_high_flag_count=0,
        llm_mode="real",
        results=results,
        decision="conditionally_ready",
        decided_at="2026-10-27T16:00:00Z",
        decision_note=(
            "The committee reviewed the manuscript and conditional readiness is granted "
            "subject to revision of the uncaptioned tables and verification of citations."
        ),
        pending_review_count=5,
        rubric_is_current=True,
        previous_status=None,
        previous_composite_score=None,
    )

    export_data = ReportExportData(
        report=report,
        flags=flags,
        flag_details=flag_details,
        page_images=page_images,
        archive_size_n=20,
        source_format="docx",
    )

    pdf_bytes = build_report_pdf(export_data)
    assert len(pdf_bytes) > 1000

    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    page_count = len(doc)
    assert page_count >= 4, f"Expected at least 4 pages, got {page_count}"

    for page_idx, page in enumerate(doc):
        page_num = page_idx + 1
        text = page.get_text()
        rect = page.rect
        blocks = page.get_text("blocks")

        # Layout bounding box check: no text should be clipped outside the page
        clipped_blocks = [
            b
            for b in blocks
            if b[0] < -0.5 or b[1] < -0.5 or b[2] > rect.width + 0.5 or b[3] > rect.height + 0.5
        ]
        assert len(clipped_blocks) == 0, (
            f"Page {page_num} has clipped text blocks: {clipped_blocks}"
        )

        # Footer page number check
        assert f"Page {page_num}" in text, (
            f"Page {page_num} missing expected footer 'Page {page_num}'"
        )

        # Running header on pages after page 1
        if page_idx > 0:
            assert "VERIDICAL - Readiness Report" in text, (
                f"Page {page_num} missing running header 'VERIDICAL - Readiness Report'"
            )

    full_text = "\n".join(p.get_text() for p in doc)
    assert "This manuscript was submitted as a DOCX" in full_text
    assert "BSIT Capstone Group 12" in full_text
    assert "Final Decision" in full_text
    assert "Decision reason" in full_text


def test_docx_reportlab_export_xml_escaping_and_watermark_invariants():
    """Verify that XML special characters (&, <, >, quotes) in flag excerpts,
    criterion text, and reasoning do not crash ReportLab mini-markup parser,
    and that DRAFT watermark appears when decision is None and disappears when decided.
    """
    # 1. Undecided report with XML special characters
    xml_hazardous_text = "Analysis of A & B < C > D 'single' \"double\" &amp; &lt; &gt;"
    results = [
        CriterionResultOut(
            check_result_id=1,
            criterion_id=1,
            text=f"Criterion with hazardous characters: {xml_hazardous_text}",
            type="structural",
            weight=1.0,
            weight_importance="high",
            kind="rule",
            outcome="failed",
            score=0.0,
            basis=f"Basis: {xml_hazardous_text}",
            anchor="paragraph 0",
            reasoning=f"Reasoning with tags <foo> & <bar>: {xml_hazardous_text}",
            reason=f"Reason: {xml_hazardous_text}",
            evidence=[],
            resolution=None,
        )
    ]

    flags = [
        FlagSummaryOut(
            id=301,
            check_kind="citations",
            severity="high",
            criterion_text=f"Integrity check: {xml_hazardous_text}",
            evidence_excerpt=f"Hazardous quote: {xml_hazardous_text}",
            page_anchor="paragraph 4",
            overridden=False,
            affects_readiness=True,
        )
    ]

    flag_details = {
        301: FlagExportDetailOut(
            flag_id=301,
            finding_headline=f"Headline: {xml_hazardous_text}",
            check_kind="citations",
            category="citation_anomaly",
            severity="high",
            page_anchor="paragraph 4",
            affects_readiness=True,
            excerpt=f"Excerpt: {xml_hazardous_text}",
            location_label="paragraph 4",
            reasoning=f"Reasoning with entities: {xml_hazardous_text}",
            override_reason=None,
        )
    }

    page_images = {
        301: PageEvidenceImageOut(
            image_png=None,
            page=None,
            reason="DOCX layout - no page image",
            reason_code="docx_layout",
        )
    }

    # Undecided report -> must have DRAFT watermark
    report_draft = ReportOut(
        check_run_id=1001,
        manuscript_id=501,
        manuscript_group_label="A" * 150,  # Very long label to stress running header
        manuscript_original_filename="draft_submission.docx",
        rubric_title="Rubric & Guidelines <2026>",
        status="not_ready",
        composite_score=45.0,
        thresholds={"ready_min_score": 85.0, "not_ready_max_score": 60.0},
        reason=f"Requires review: {xml_hazardous_text}",
        flag_deduction=15.0,
        unresolved_high_flag_count=1,
        llm_mode="real",
        results=results,
        decision=None,
        decided_at=None,
        decision_note=None,
        pending_review_count=1,
        rubric_is_current=True,
        previous_status=None,
        previous_composite_score=None,
    )

    export_draft = ReportExportData(
        report=report_draft,
        flags=flags,
        flag_details=flag_details,
        page_images=page_images,
        archive_size_n=5,
        source_format="docx",
    )

    pdf_bytes_draft = build_report_pdf(export_draft)
    doc_draft = fitz.open(stream=pdf_bytes_draft, filetype="pdf")
    text_draft = "\n".join(p.get_text() for p in doc_draft)
    assert "DRAFT" in text_draft

    # 2. Decided report -> must NOT have DRAFT watermark
    from datetime import UTC, datetime

    report_decided = report_draft.model_copy(
        update={
            "decision": "not_ready",
            "decided_at": datetime(2026, 10, 27, 17, 0, 0, tzinfo=UTC),
            "decision_note": f"Official faculty decision: {xml_hazardous_text}",
        }
    )
    export_decided = export_draft.model_copy(update={"report": report_decided})

    pdf_bytes_decided = build_report_pdf(export_decided)
    doc_decided = fitz.open(stream=pdf_bytes_decided, filetype="pdf")
    text_decided = "\n".join(p.get_text() for p in doc_decided)
    # The large DRAFT watermark text should not be present in text layer of decided report
    assert "DRAFT" not in text_decided


def test_manuscript_file_format_more_adversarial_paths():
    """Test Windows backslashes, mixed slashes, whitespace, and empty nodes."""
    # 1. Windows path with backslashes
    m_win = Manuscript(
        id=91,
        original_filename="C:\\Users\\Student\\Documents\\Capstone_Report.docx",
        file_ref="C:\\Veridical\\data\\uploads\\blob_91.docx",
    )
    assert m_win.file_format == "docx"

    # 2. Filename with internal spaces
    m_spaces_clean = Manuscript(
        id=93,
        original_filename="manuscript with internal spaces.docx",
        file_ref=None,
    )
    assert m_spaces_clean.file_format == "docx"

    # 3. Empty SectionTree with source='heuristics' falls back to pdf if no other signal
    m_empty_heuristics = Manuscript(
        id=94,
        original_filename="raw_file_no_ext",
        file_ref="remote/key",
        section_tree={"source": "heuristics", "nodes": []},
    )
    assert m_empty_heuristics.file_format == "pdf"

    # 4. SectionTree with malformed or non-dict nodes list does not crash
    m_malformed_nodes = Manuscript(
        id=95,
        original_filename="malformed_doc",
        file_ref="remote/key",
        section_tree={"source": "heuristics", "nodes": [None, "string_node", 123]},
    )
    assert m_malformed_nodes.file_format == "pdf"

